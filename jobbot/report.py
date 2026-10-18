"""Date-wise CSV files that open in Excel: every job scraped that day with its email and status.

reports/<date>/jobs_<date>.csv    every job found that day (kept, skipped, matched, sent ...)
reports/<date>/emails_<date>.csv  every email sent that day, whenever the job was found
Files are rewritten on each run, so the status columns are always current.
"""
import csv
import json

from . import HomePath

REPORTS = HomePath("reports")

JOB_COLUMNS = [
    ("Date found", lambda r: (r["found_at"] or "")[:10]), ("Job ID", lambda r: r["id"]),
    ("Source", lambda r: r["source"]), ("Job title", lambda r: r["title"]), ("Company", lambda r: r["company"]),
    ("Country / region", lambda r: "Taiwan" if r["region"] == "taiwan" else "Remote / abroad"),
    ("Location", lambda r: r["location"]),
    ("Job type", lambda r: {"full": "Full-time", "part": "Part-time"}.get(r["job_type"], r["job_type"])),
    ("Email track", lambda r: r["track"]), ("Field", lambda r: r["category"]), ("Match score", lambda r: r["score"]),
    ("Stage", lambda r: stage(r)), ("Reason skipped", lambda r: r["skip_reason"] or ""),
    ("HR name", lambda r: r["hr_name"] or ""), ("HR email", lambda r: ", ".join(json.loads(r["emails"] or "[]"))),
    ("Email found where", lambda r: (r["email_source"] or ("in the job post" if r["emails"] not in (None, "[]") else ""))),
    ("Resume written by", lambda r: r["tailored_by"] or ""), ("Resume file", lambda r: r["resume_path"] or ""),
    ("Email subject", lambda r: r["email_subject"] or ""),
    ("Sent from", lambda r: r["sender"] or ""), ("Sent to", lambda r: r["to_email"] or ""),
    ("Sent at", lambda r: r["sent_at"] or ""), ("Result", lambda r: r["outcome"] or ""),
    ("Send error", lambda r: r["send_error"] or ""),
    ("Deadline (leads)", lambda r: r["deadline"] or ""),
    ("Salary", lambda r: r["salary"] or ""), ("Posted on", lambda r: r["posted_at"] or ""), ("Job link", lambda r: r["url"]),
]


def stage(r):
    """One plain word for where the job stands."""
    if r["outcome"]:
        return {"replied": "Reply received", "bounced": "Email bounced",
                "bad address": "Bad address"}.get(r["outcome"], r["outcome"].capitalize())
    if r["send_status"] == "sent":
        return "Email sent"
    if r["send_status"] == "sending":
        return "Interrupted while sending: check Sent folder"
    if r["send_status"] == "failed":
        return "Send failed"
    if r["send_status"] == "rejected":
        return "Rejected by you"
    if r["send_status"] == "approved":
        return "Approved, waiting to send"
    if r["status"] == "skipped":
        return "Skipped by filter"
    if r["status"] == "new":
        return "Kept, not matched yet"
    if r["source"] == "freelancer":
        return "Bid on the platform"
    if r["emails"] in (None, "[]"):
        return "Matched, no HR email: apply on the site"
    if not r["resume_path"] and r["job_type"] != "lead":
        return "Matched, resume not written yet"
    return "Ready, waiting for your review"


def _write(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:   # utf-8-sig so Excel shows Chinese correctly
        w = csv.writer(f)
        w.writerow([name for name, _ in JOB_COLUMNS])
        for r in rows:
            w.writerow([get(r) for _, get in JOB_COLUMNS])
    return path


def write_day(con, day):
    """Writes both files for one date. Returns (jobs_path, n_jobs, emails_path, n_emails)."""
    jobs = con.execute("SELECT * FROM jobs WHERE found_at LIKE ? ORDER BY (status='matched') DESC, score DESC",
                       (day + "%",)).fetchall()
    sent = con.execute("SELECT * FROM jobs WHERE sent_at LIKE ? ORDER BY sent_at", (day + "%",)).fetchall()
    folder = REPORTS / day
    return (_write(folder / f"jobs_{day}.csv", jobs), len(jobs),
            _write(folder / f"emails_{day}.csv", sent), len(sent))


def days_with_data(con):
    return [r[0] for r in con.execute(
        "SELECT DISTINCT substr(found_at,1,10) FROM jobs UNION SELECT DISTINCT substr(sent_at,1,10) FROM jobs "
        "WHERE sent_at IS NOT NULL ORDER BY 1") if r[0]]


# ---------- one running tracker of every application that was sent ----------

TRACKER = HomePath("reports", "Applications-Sent.xlsx")
TRACK_NAMES = {"full": "Full-time job", "part": "Part-time job", "freelance": "Freelance work", "lead": "Lead",
               "collaboration": "Collaboration", "subcontract": "Subcontracting lead"}


def delivery(r):
    if r["outcome"] == "bounced":
        return "Bounced (address did not accept it)"
    if r["outcome"] == "replied":
        return "Reply received"
    if r["auto_ack"]:
        return "Delivered (automatic acknowledgement received)"
    return "Sent, no bounce"


TRACKER_COLUMNS = [
    ("No.", 6, None), ("Sent on", 12, lambda r: r["sent_at"][:10]), ("Time", 9, lambda r: r["sent_at"][11:16]),
    ("Kind", 15, lambda r: TRACK_NAMES.get(r["track"], r["track"])), ("Field", 18, lambda r: r["category"]),
    ("Company", 30, lambda r: r["company"]), ("Role applied for", 44, lambda r: r["title"]),
    ("Country / region", 16, lambda r: "Taiwan" if r["region"] == "taiwan" else "Remote / abroad"),
    ("Sent to (HR email)", 32, lambda r: r["to_email"]), ("Contact name", 16, lambda r: r["hr_name"] or ""),
    ("Sent from", 30, lambda r: r["sender"]), ("CC", 26, lambda r: r["cc_email"] or ""),
    ("Status", 40, delivery), ("Status updated", 17, lambda r: (r["outcome_at"] or "")[:16].replace("T", " ")),
    ("Follow-up sent", 17, lambda r: (r["followup_at"] or "")[:16].replace("T", " ")),
    ("Email subject", 50, lambda r: r["email_subject"]), ("Resume sent", 46, lambda r: (r["resume_path"] or "").split("\\")[-1]),
    ("Resume written by", 16, lambda r: {"ai": "AI", "basic": "Template", "template": "Template"}.get(r["tailored_by"], "")),
    ("Job source", 15, lambda r: r["source"]), ("Job link", 50, lambda r: r["url"]),
    ("Match score", 11, lambda r: r["score"]), ("Job ID", 8, lambda r: r["id"]),
]


def write_tracker(con):
    """Rewrite the tracker from the database: one row per email that was sent, oldest first.
    Also writes a .csv copy. Returns (path, rows). If the Excel file is open, a copy is written."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    rows = con.execute("SELECT * FROM jobs WHERE send_status='sent' ORDER BY sent_at").fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications sent"
    ws.append([name for name, _, _ in TRACKER_COLUMNS])
    for n, r in enumerate(rows, 1):
        ws.append([n if get is None else get(r) for _, _, get in TRACKER_COLUMNS])
    for i, (name, width, _) in enumerate(TRACKER_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center")
    link_col = [name for name, _, _ in TRACKER_COLUMNS].index("Job link") + 1
    for row in range(2, len(rows) + 2):
        cell = ws.cell(row=row, column=link_col)
        if str(cell.value or "").startswith("http"):
            cell.hyperlink, cell.font = cell.value, Font(color="0563C1", underline="single")
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = ws.dimensions
    summary = wb.create_sheet("Summary")
    summary.append(["Total applications sent", len(rows)])
    for label, where in (("Replies received", "outcome='replied'"), ("Automatic acknowledgements", "auto_ack=1"),
                         ("Bounced", "outcome='bounced'")):
        summary.append([label, con.execute(f"SELECT COUNT(*) FROM jobs WHERE send_status='sent' AND {where}").fetchone()[0]])
    summary.append([])
    summary.append(["By day", "Sent"])
    for r in con.execute("SELECT substr(sent_at,1,10) d, COUNT(*) c FROM jobs WHERE send_status='sent' GROUP BY 1 ORDER BY 1"):
        summary.append([r["d"], r["c"]])
    summary.append([])
    summary.append(["By kind", "Sent"])
    for r in con.execute("SELECT track, COUNT(*) c FROM jobs WHERE send_status='sent' GROUP BY 1 ORDER BY 2 DESC"):
        summary.append([TRACK_NAMES.get(r["track"], r["track"]), r["c"]])
    summary.column_dimensions["A"].width, summary.column_dimensions["B"].width = 32, 12
    for cell in summary["A"]:
        cell.font = Font(bold=True)
    REPORTS.mkdir(exist_ok=True)
    path = TRACKER
    try:
        wb.save(path)
    except PermissionError:   # the file is open in Excel
        path = REPORTS / "Applications-Sent (new copy - close Excel).xlsx"
        wb.save(path)
    with open(REPORTS / "Applications-Sent.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow([name for name, _, _ in TRACKER_COLUMNS])
        for n, r in enumerate(rows, 1):
            w.writerow([n if get is None else get(r) for _, _, get in TRACKER_COLUMNS])
    return path, len(rows)


# ---------- Excel list of every job from the last two weeks ----------

JOBS_BOOK = HomePath("reports", "Jobs-Last-14-Days.xlsx")
TYPE_NAMES = {"full": "Full-time", "part": "Part-time", "freelance": "Freelance", "lead": "Lead (tender)",
              "collaboration": "Collaboration", "subcontract": "Subcontracting lead"}


def _place(r):
    """(country, city) from the stored location text."""
    loc = (r["location"] or "").strip()
    if r["region"] == "taiwan":
        return "Taiwan", loc.replace("Taiwan", "").strip(" ,")
    parts = [x.strip() for x in loc.replace(";", ",").split(",") if x.strip()]
    if not parts:
        return ("Remote" if r["remote"] else ""), ""
    return (parts[-1], ", ".join(parts[:-1])) if len(parts) > 1 else (parts[0], "")


JOB_BOOK_COLUMNS = [
    ("Position", 44, lambda r: r["title"]), ("Company", 30, lambda r: r["company"]),
    ("Type", 13, lambda r: TYPE_NAMES.get(r["track"] or r["job_type"], r["job_type"])),
    ("Field", 16, lambda r: r["category"]),
    ("Published", 12, lambda r: r["posted_at"] or ""), ("Found by bot", 12, lambda r: (r["found_at"] or "")[:10]),
    ("Country", 16, lambda r: _place(r)[0]), ("City / town", 22, lambda r: _place(r)[1]),
    ("Remote", 8, lambda r: "Yes" if r["remote"] else ""),
    ("HR email", 32, lambda r: ", ".join(json.loads(r["emails"] or "[]"))),
    ("Contact name", 14, lambda r: r["hr_name"] or ""),
    ("Job link", 46, lambda r: r["url"]), ("Company website", 30, lambda r: r["company_url"] or ""),
    ("Stage", 34, stage), ("Match score", 11, lambda r: r["score"]), ("Salary", 22, lambda r: r["salary"] or ""),
    ("Platform", 16, lambda r: r["source"]),
    ("Email found where", 30, lambda r: r["email_source"] or ("in the job post" if r["emails"] not in (None, "[]") else "")),
    ("Sent to", 28, lambda r: r["to_email"] or ""), ("Sent at", 16, lambda r: (r["sent_at"] or "").replace("T", " ")[:16]),
    ("Job description", 90, lambda r: (r["description"] or "")[:6000]),
    ("Job ID", 8, lambda r: r["id"]),
]


def write_jobs_book(con, days=14):
    """All jobs published (or, when the site gives no date, found) in the last `days` days.
    Sheets: all kept jobs, the ones with an HR email, and the ones the filter left out."""
    import datetime as dt
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    since = (dt.date.today() - dt.timedelta(days=days)).isoformat()
    recent = "COALESCE(NULLIF(posted_at, ''), substr(found_at, 1, 10)) >= ?"
    order = " ORDER BY COALESCE(NULLIF(posted_at, ''), substr(found_at, 1, 10)) DESC, score DESC"
    kept = con.execute(f"SELECT * FROM jobs WHERE status!='skipped' AND {recent}" + order, (since,)).fetchall()
    out = con.execute(f"SELECT * FROM jobs WHERE status='skipped' AND {recent}" + order, (since,)).fetchall()
    wb = Workbook()

    def sheet(ws, rows, columns):
        ws.append([name for name, _, _ in columns])
        for r in rows:
            ws.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in (get(r) for _, _, get in columns)])
        names = [name for name, _, _ in columns]
        for i, (name, width, _) in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
            c = ws.cell(row=1, column=i)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
        for col in ("Job link", "Company website"):
            if col in names:
                idx = names.index(col) + 1
                for row in range(2, len(rows) + 2):
                    c = ws.cell(row=row, column=idx)
                    if str(c.value or "").startswith("http"):
                        c.hyperlink, c.font = c.value, Font(color="0563C1", underline="single")
        if "Job description" in names:
            idx = names.index("Job description") + 1
            for row in range(2, len(rows) + 2):
                ws.cell(row=row, column=idx).alignment = Alignment(wrap_text=False, vertical="top")
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "All jobs"
    sheet(ws, kept, JOB_BOOK_COLUMNS)
    sheet(wb.create_sheet("With HR email"), [r for r in kept if r["emails"] not in (None, "[]")], JOB_BOOK_COLUMNS)
    slim = [c for c in JOB_BOOK_COLUMNS if c[0] not in ("Job description", "Sent to", "Sent at", "Email found where")]
    slim.insert(4, ("Why left out", 30, lambda r: r["skip_reason"] or ""))
    sheet(wb.create_sheet("Left out by filter"), out, slim)
    info = wb.create_sheet("Summary")
    info.append([f"Jobs published in the last {days} days", f"(since {since})"])
    info.append(["Kept by the filter", len(kept)])
    info.append(["  of which with an HR email", sum(r["emails"] not in (None, "[]") for r in kept)])
    info.append(["Left out by the filter", len(out)])
    info.append([])
    info.append(["By type", "Jobs", "With email"])
    for key, label in TYPE_NAMES.items():
        rows = [r for r in kept if (r["track"] or r["job_type"]) == key]
        info.append([label, len(rows), sum(r["emails"] not in (None, "[]") for r in rows)])
    info.append([])
    info.append(["By country", "Jobs"])
    counts = {}
    for r in kept:
        counts[_place(r)[0] or "Not stated"] = counts.get(_place(r)[0] or "Not stated", 0) + 1
    for country, n in sorted(counts.items(), key=lambda x: -x[1])[:40]:
        info.append([country, n])
    info.column_dimensions["A"].width = 38
    for c in info["A"]:
        c.font = Font(bold=True)
    REPORTS.mkdir(exist_ok=True)
    path = JOBS_BOOK
    try:
        wb.save(path)
    except PermissionError:   # open in Excel
        path = REPORTS / "Jobs-Last-14-Days (new copy - close Excel).xlsx"
        wb.save(path)
    return path, len(kept), sum(r["emails"] not in (None, "[]") for r in kept)
