"""Feature tests, one by one. Run from the jobbot folder:   python tests\\run_all.py

Safe to run at any time: nothing is sent to an HR address. One test email goes from your sending
account to your own other address. Test rows written to the database are removed again.
Results are printed and written to TEST-REPORT.md next to config.yaml.

To run only some areas, name them:   python tests\\run_all.py sources api
(the report file is then left as it is).
"""
import datetime as dt
import json
import os
import re
import sqlite3
import sys
import threading
import time
import urllib.parse
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))


def _test_data():
    """Run against a copy of the live data in tests/.data, so nothing the tests do can reach the live database.
    Set JOBBOT_TEST_LIVE=1 to test the live data itself (the old behaviour)."""
    if os.environ.get("JOBBOT_TEST_LIVE") == "1" or os.environ.get("JOBBOT_DATA"):
        return None
    live = ROOT
    data = ROOT / "tests" / ".data"
    import shutil
    shutil.rmtree(data, ignore_errors=True)
    data.mkdir(parents=True)
    src = sqlite3.connect(live / "jobbot.db")
    dst = sqlite3.connect(data / "jobbot.db")
    src.backup(dst)
    src.close(), dst.close()
    for f in [live / ".env", *live.glob("accounts/*/.env")]:
        target = data / f.relative_to(live)
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(f, target)
    os.environ["JOBBOT_DATA"] = str(data)
    return data


TEST_DATA = _test_data()

from jobbot import api, cleanup, contacts as ct, db, discover as dc, filter as flt, leads, mailer, render  # noqa: E402
from jobbot import control, profilestore, report as rp, scheduler, sendqueue as sq, sourcestore  # noqa: E402
from jobbot import DATA, tailor as tl, use_home  # noqa: E402
from jobbot.config import (_leaves, enabled, file_config, load_config, load_sources, overrides, reset_setting,  # noqa: E402
                           set_setting, setting_problem, settings_list)
from jobbot.sources import ADAPTERS, generic  # noqa: E402
from jobbot.sources.base import Job  # noqa: E402

RESULTS = []
ONLY = {a.lower() for a in sys.argv[1:]}


def test(area, name):
    def wrap(fn):
        if ONLY and area.lower() not in ONLY:
            return fn
        start = time.time()
        try:
            detail = fn() or ""
            state = "SKIP" if str(detail).startswith("SKIP") else "PASS"
        except AssertionError as e:
            state, detail = "FAIL", str(e) or "assertion failed"
        except Exception as e:
            state, detail = "FAIL", f"{type(e).__name__}: {str(e)[:160]}"
        RESULTS.append((area, name, state, str(detail)[:170], time.time() - start))
        print(f"{state:<5} {area:<12} {name:<52} {str(detail)[:90]}", flush=True)
        return fn
    return wrap


cfg = load_config()
master = render.load_master()
con = db.connect()


def job(**kw):
    base = dict(source="test", ext_id="t1", title="Senior .NET Developer", company="Acme", url="https://example.org/j",
                description="C#, .NET 8, ASP.NET Core, SQL Server, Angular. Fully remote.", location="Worldwide",
                region="abroad", remote=True, posted_at=dt.date.today().isoformat())
    base.update(kw)
    return Job(**base)


# ---------------------------------------------------------------- settings
@test("Settings", "the database structure is at the latest step")
def _():
    # one database (db.py): shared tables, per-account tables, the control tables, and the account's views
    versions = dict(con.execute("SELECT name, version FROM schema_versions").fetchall())
    assert versions == {"jobs": len(db.SCHEMA_STEPS), "control": len(control.STEPS)}, versions
    tables = {r[0] for r in con.execute("SELECT name FROM main.sqlite_master WHERE type='table'")}
    need = {"postings", "applications", "acct_settings", "acct_documents", "acct_prefs", "acct_process_log", "acct_meta", "runs",
            "sources", "candidates", "countries", "catalogue", "accounts", "users", "plans", "payments"}
    assert need <= tables, f"missing tables: {need - tables}"
    views = {r[0] for r in con.execute("SELECT name FROM temp.sqlite_master WHERE type='view'")}
    assert set(db.VIEWS) <= views, f"missing views: {set(db.VIEWS) - views}"
    assert not tables & set(db.VIEWS), f"old per-account tables still in the file: {tables & set(db.VIEWS)}"
    return f"one file, {len(tables)} tables, views {', '.join(sorted(views))}"


@test("Settings", "every setting is a row in the settings table, with its default")
def _():
    rows = {r["key"]: r for r in settings_list()}
    defaults = dict(_leaves(file_config()))
    assert set(defaults) <= set(rows), f"not in the table: {set(defaults) - set(rows)}"
    assert all(rows[k]["default"] == v for k, v in defaults.items()), "a default differs from config.yaml"
    changed = overrides()
    assert all(rows[k]["value"] == v for k, v in defaults.items() if k not in changed)
    return f"{len(rows)} settings, {len(changed)} changed from the default"


@test("Settings", "a value of the wrong kind, or for an unknown setting, is refused")
def _():
    assert setting_problem("sending.per_run", 5) is None and setting_problem("process.leads", False) is None
    for key, value in (("sending.per_run", "five"), ("sending.review", 3), ("regions", "Taiwan"), ("no.such", 1),
                       ("process.nope", True), ("process.leads", "off"), ("sending.per_run", None)):
        assert setting_problem(key, value), f"{key}={value!r} was accepted"


@test("Settings", "settings load with all sections")
def _():
    for key in ("regions", "daily_target", "min_score", "search", "profile", "sending", "ai", "leads", "sources", "loop"):
        assert key in cfg, f"missing section {key}"
    return f"{len(cfg)} sections"


@test("Settings", "a changed setting is used, and reset goes back to the default")
def _():
    before = load_config()["sending"]["per_run"]
    set_setting("sending.per_run", before + 5)
    try:
        assert load_config()["sending"]["per_run"] == before + 5
    finally:
        reset_setting("sending.per_run")
    assert load_config()["sending"]["per_run"] == before
    return f"{before} -> {before + 5} -> {before}"


@test("Settings", "a process can be switched off and on")
def _():
    set_setting("process.leads", False)
    try:
        assert not enabled(load_config(), "leads")
    finally:
        reset_setting("process.leads")
    assert enabled(load_config(), "leads")


# ---------------------------------------------------------------- sources
@test("Sources", "every source in the sources table has a reader")
def _():
    srcs = load_sources()
    assert len(srcs) == con.execute("SELECT COUNT(*) FROM sources").fetchone()[0] > 0
    missing = [s["name"] for s in srcs if s["adapter"] not in ADAPTERS]
    assert not missing, f"no reader for {missing}"
    return f"{len(srcs)} sources, {sum(bool(s.get('enabled')) for s in srcs)} switched on"


@test("Sources", "a source is saved, renamed on a clash, changed and removed in the table")
def _():
    first = dc.add_source("zz-test-source", "https://example.org/jobs", "Remote", {"kind": "feed", "url": "https://example.org/feed"}, False)
    second = dc.add_source("zz-test-source", "https://example.org/jobs2", "Remote", {"kind": "feed", "url": "https://example.org/feed2"}, False)
    try:
        assert (first, second) == ("zz-test-source", "zz-test-source-2")
        assert sourcestore.get(second)["params"]["name"] == second, "jobs would be labelled with the wrong source name"
        assert sourcestore.set_enabled(first, True) and sourcestore.get(first)["enabled"]
        assert sourcestore.update(first, {"region": "Japan"}) and sourcestore.get(first)["params"]["region"] == "Japan"
        assert sourcestore.remove("104") == "built-in" and sourcestore.get("104"), "a built-in source was deleted"
    finally:
        removed = [sourcestore.remove(first), sourcestore.remove(second)]
    assert removed == ["removed", "removed"] and not sourcestore.get(first)
    assert sourcestore.remove(first) == "missing" and not sourcestore.set_enabled(first, True)


@test("Sources", "candidate sites and their test results are in the candidates table")
def _():
    rows = sourcestore.candidates()
    readable = [c for c in rows if c["state"] == "readable"]
    assert rows and readable, "no candidates, or none readable"
    assert all(c["url"].startswith("http") and c["region"] for c in rows)
    linked = [c for c in readable if c["source"]]
    return f"{len(rows)} sites, {len(readable)} readable, {len(linked)} of those are sources"


@test("Sources", "bidding-platform projects are listed but never emailed")
def _():
    bid = Job(source="freelancer", ext_id="zz-test-bid", title="Build a website", company="Freelancer.com client", url="u",
              description="client@bid-test.example", location="Remote", region="", remote=True, job_type="part", posted_at="")
    try:
        db.insert_job(con, bid, {"track": "freelance", "category": "x", "emails": ["client@bid-test.example"], "score": 90,
                                 "notes": "", "skip": None, "dedupe_key": "bidtest|x"})
        con.execute("UPDATE jobs SET status='matched', send_status='approved', resume_path='x' WHERE source='freelancer' AND ext_id='zz-test-bid'")
        assert con.execute("SELECT 1 FROM jobs WHERE ext_id='zz-test-bid' AND send_status='approved'").fetchone()
        assert "client@bid-test.example" not in [to for _, to in sq.sendable(con, cfg)]
        try:
            con.execute("UPDATE jobs SET source='other' WHERE ext_id='zz-test-bid'")
            raise AssertionError("a posting's source could be changed")
        except sqlite3.IntegrityError:
            pass   # source and ext_id identify a shared posting: changing them is refused
    finally:
        con.rollback()


def _read(name):
    s = next(s for s in load_sources() if str(s["name"]) == name)
    params = dict(s.get("params") or {})
    if s["adapter"] == "generic":
        params.setdefault("name", name); params.setdefault("region", s.get("region", "Remote"))
    # a sample is enough to show the reader works: two keywords, one page each
    small = cfg | {"search": cfg["search"] | {"keywords": cfg["search"]["keywords"][:2],
                                              "part_time_keywords": cfg["search"].get("part_time_keywords", [])[:1]}}
    params.update({k: 1 for k in ("pages", "remote_pages") if k in params or s["adapter"] == "tw104"})
    return ADAPTERS[s["adapter"]](small, params)


for _name in ("104", "remotive", "weworkremotely", "hn-whoishiring", "ted-tenders"):
    @test("Sources", f"live read: {_name}")
    def _(n=_name):
        jobs = _read(n)
        assert len(jobs) > 0, "no jobs returned"
        assert all(j.title and j.url for j in jobs[:20]), "a job without title or link"
        return f"{len(jobs)} items"


@test("Sources", "generic reader recognises an RSS feed")
def _():
    params, jobs = generic.detect("https://weworkremotely.com/remote-jobs.rss", "t", "Remote")
    assert params and params["kind"] == "feed" and jobs
    return f"{len(jobs)} jobs"


@test("Sources", "generic reader recognises a hiring-system board")
def _():
    params, jobs = generic.detect("https://boards.greenhouse.io/gitlab", "t", "Remote")
    assert params and params["kind"] == "greenhouse" and jobs
    return f"greenhouse, {len(jobs)} jobs"


@test("Sources", "generic reader refuses a page with no jobs")
def _():
    params, jobs = generic.detect("https://example.org/", "t", "Remote")
    assert params is None and not jobs


@test("Discovery", "country list covers the world")
def _():
    n = len(dc.all_countries())
    assert n >= 190, n
    return f"{n} countries and regions"


@test("Discovery", "rotation picks countries not checked yet")
def _():
    nxt = dc.next_countries(con, 3)
    done = {r["region"] for r in con.execute("SELECT region FROM discovery")}
    assert nxt and not (set(nxt) & done), f"{nxt} overlaps checked ones"
    return f"next: {', '.join(nxt)}; checked so far: {len(done)}"


@test("Discovery", "catalogue offers candidates without a web search")
def _():
    c = dc.candidates("Remote", use_search=False)
    assert len(c) >= 5
    return f"{len(c)} candidate sites for Remote"


# ---------------------------------------------------------------- filter
@test("Filter", "a strong remote .NET job scores high and is kept")
def _():
    v = flt.judge(job(), cfg)
    assert not v["skip"] and v["score"] >= 60 and v["category"] == "Software .NET", v
    return f"score {v['score']}, field {v['category']}"


@test("Filter", "on-site job abroad: kept when open to relocation, skipped when not")
def _():
    onsite = dict(remote=False, location="Berlin, Germany")
    assert flt.judge(job(**onsite), {**cfg, "open_to_relocation": False})["skip"] == "on-site outside your regions"
    assert flt.judge(job(**onsite), {**cfg, "open_to_relocation": True})["skip"] is None
    return f"current setting: open_to_relocation = {cfg.get('open_to_relocation')}"


@test("Filter", "Mandarin-only post is skipped")
def _():
    v = flt.judge(job(title="資深後端工程師", description="負責系統開發與維護，需具備良好溝通能力", region="taiwan", remote=False), cfg)
    assert v["skip"] == "post mostly in Chinese", v["skip"]


@test("Filter", "remote-only countries: remote and freelance kept, on-site left out")
def _():
    rules = {"India": ["india"]}
    c = cfg | {"remote_only_countries": rules}
    assert "only remote or freelance" in flt.judge(job(remote=False, location="Pune, India"), c)["skip"]
    assert flt.judge(job(remote=True, location="India"), c)["skip"] is None
    assert flt.judge(job(remote=False, location="Mumbai, India", track="freelance"), c)["skip"] is None
    live = sorted(cfg.get("remote_only_countries") or {})
    return f"current setting: remote only in {', '.join(live) or 'no country'}"


@test("Filter", "Taiwan: technical any type, non-technical only part-time or temporary")
def _():
    tw = dict(region="taiwan", remote=False, location="Taipei, Taiwan", description="English speaking team.")
    assert flt.judge(job(title="Backend Engineer", job_type="full", **tw), cfg)["skip"] is None
    assert "non-technical full-time" in flt.judge(job(title="Marketing Specialist", job_type="full", **tw), cfg)["skip"]
    assert flt.judge(job(title="Marketing Assistant", job_type="part", **tw), cfg)["skip"] is None


@test("Filter", "post older than two weeks is skipped")
def _():
    old = (dt.date.today() - dt.timedelta(days=cfg["max_age_days"] + 3)).isoformat()
    assert flt.judge(job(posted_at=old), cfg)["skip"] == "too old"


@test("Filter", "part-time and freelance are recognised")
def _():
    a, b = job(title="Part-time C# Developer"), job(title="Freelance .NET Developer")
    va, vb = flt.judge(a, cfg), flt.judge(b, cfg)
    assert a.job_type == "part" and va["track"] == "part", (a.job_type, va["track"])
    assert b.job_type == "part" and vb["track"] == "freelance", (b.job_type, vb["track"])


@test("Filter", "the same job from the same source is stored once")
def _():
    def clear():
        con.execute("DELETE FROM jobs WHERE source='test'")
        con.execute("DELETE FROM postings WHERE source='test'"); con.commit()
    clear()
    j = job(ext_id="dup-test-1")
    v = flt.judge(j, cfg)
    try:
        first, second = db.insert_job(con, j, v), db.insert_job(con, j, v)
    finally:
        clear()
    assert first and not second


# ---------------------------------------------------------------- emails found
@test("HR emails", "non-hiring and placeholder addresses are dropped")
def _():
    bad = ["privacy@gmail.com", "reasonable-accommodations@gmail.com", "email@gmail.com", "support@gmail.com",
           "noreply@gmail.com", "accessibility@gmail.com"]
    kept = ct.clean(bad + ["jobs@gmail.com"])
    assert kept == ["jobs@gmail.com"], kept


@test("HR emails", "an address on a domain without a mail server is dropped")
def _():
    assert ct.clean(["hr@this-domain-does-not-exist-9x7q.example"]) == []


@test("HR emails", "hiring addresses rank before general inboxes")
def _():
    assert ct.clean(["info@gmail.com", "careers@gmail.com"])[0] == "careers@gmail.com"


# ---------------------------------------------------------------- resume
@test("Resume", "master resume loads with the AMG end date")
def _():
    assert master["experience"][0]["dates"].startswith("Mar 2025 – Sep 2026"), master["experience"][0]["dates"]
    return f"{len(master['experience'])} positions"


@test("Resume", "Gold Card appears for Taiwan, never for abroad")
def _():
    tw, ab = render.to_html(render.build("taiwan")), render.to_html(render.build("abroad"))
    assert tw.lower().count("gold card") == 2 and "gold card" not in ab.lower()


@test("Resume", "a Gold Card line in a non-Taiwan resume is blocked")
def _():
    try:
        render.check_region_rule(render.to_html(render.build("taiwan")), "abroad")
    except ValueError:
        return "blocked"
    raise AssertionError("not blocked")


@test("Resume", "PDF is created in Chrome: 2 pages, no tool name in properties")
def _():
    import pymupdf
    out = render.OUTPUT / "test" / "feature_test.pdf"
    render.render("abroad", out, channel=cfg["browser"]["channel"])
    d = pymupdf.open(out)
    meta = d.metadata
    assert len(d) == 2, f"{len(d)} pages"
    assert not meta.get("creator") and not meta.get("producer"), meta
    assert meta.get("author") == master["name"]
    return f"{out.stat().st_size // 1024} KB, author {meta['author']}"


@test("Resume", "fact check catches an invented skill")
def _():
    t = {"headline": master["headline"], "summary": "Engineer with deep Kubernetes and Rust experience.",
         "skills": master["skills"], "experience": [{"bullets": e["bullets"]} for e in master["experience"]], "email_body": "Hello."}
    row = {"title": "X", "company": "Y", "region": "abroad", "hr_name": ""}
    bad = tl.violations(t, master, row)
    assert "Kubernetes" in bad and "Rust" in bad, bad
    return f"caught: {', '.join(bad)}"


@test("Resume", "fact check passes text that only uses real facts")
def _():
    t = {"headline": master["headline"], "summary": master["summary"], "skills": master["skills"],
         "experience": [{"bullets": e["bullets"]} for e in master["experience"]], "email_body": "I work with C# and SQL Server."}
    bad = tl.violations(t, master, {"title": "X", "company": "Y", "region": "abroad", "hr_name": ""})
    assert not bad, bad


def _row(i):
    return con.execute("SELECT * FROM jobs WHERE id=?", (i,)).fetchone()


@test("Resume", "template email: role in subject, Gold Card only for Taiwan")
def _():
    tw = con.execute("SELECT * FROM jobs WHERE region='taiwan' AND job_type IN ('full','part') AND status!='skipped' LIMIT 1").fetchone()
    ab = con.execute("SELECT * FROM jobs WHERE region='abroad' AND job_type='full' AND status!='skipped' LIMIT 1").fetchone()
    a, b = tl.basic(master, tw), tl.basic(master, ab)
    if tw["track"] != "freelance":
        assert "gold card" in a["email_body"].lower()
    assert "gold card" not in b["email_body"].lower()
    assert "{" not in a["email_body"] + b["email_body"] + a["email_subject"] + b["email_subject"]
    return b["email_subject"][:70]


@test("Resume", "AI writes a job-specific resume and email")
def _():
    chain, notes = tl.engines(cfg)
    if not chain:
        return "SKIP no working AI key"
    r = con.execute("SELECT * FROM jobs WHERE source='remoteok' AND status!='skipped' AND length(description)>800 ORDER BY score DESC LIMIT 1").fetchone()
    t = tl.tailor(r, cfg, master, chain)
    if t["mode"] != "ai":
        return f"SKIP AI not available right now ({t['note'][:80]})"
    # check what the AI wrote; the signature and the "based in Taiwan (UTC+8)" line are added by the bot afterwards
    bad = tl.violations({**t, "email_body": ""}, master, r)
    assert not bad, f"not in master resume: {bad}"
    assert "gold card" not in (t["email_body"] + json.dumps(t["skills"])).lower(), "Gold Card in a non-Taiwan application"
    return f"written by {t['model']}; subject: {t['email_subject'][:50]}"


@test("Resume", "on-site job abroad: email says open to relocating, no Gold Card")
def _():
    r = dict(con.execute("SELECT * FROM jobs WHERE region='abroad' AND status!='skipped' AND job_type='full' LIMIT 1").fetchone())
    r["remote"] = 0
    body = tl.basic(master, r)["email_body"].lower()
    assert "open to relocating" in body and "remote work across time zones" not in body and "gold card" not in body
    r["remote"] = 1
    assert "remote work across time zones" in tl.basic(master, r)["email_body"].lower()


@test("Interview", "interview prep exists for a job, round by round")
def _():
    r = con.execute("SELECT * FROM jobs WHERE interview_prep IS NOT NULL ORDER BY interview_prep_at DESC LIMIT 1").fetchone()
    assert r, "no guide written yet (run: jobbot prep <id>)"
    text = r["interview_prep"]
    rounds = text.count("## Round")
    assert rounds >= 3 and "Gaps to close" in text and "Questions to ask" in text, f"{rounds} rounds"
    if r["region"] != "taiwan":
        assert "gold card" not in text.lower()
    return f"job #{r['id']}: {rounds} rounds, {len(text.split())} words"


# ---------------------------------------------------------------- sending rules
ACCOUNTS = [s for s in mailer.senders() if s["can_send"]]
ALLOWED = [e.lower() for e in cfg["sending"].get("send_from") or []]
SENDING = [s for s in ACCOUNTS if not ALLOWED or s["email"].lower() in ALLOWED]


def planned(c2):
    # lift the daily limit for the test only, so "today's quota is used up" cannot hide the result
    c2["sending"]["warmup_caps"], c2["sending"]["daily_ceiling"], c2["sending"]["per_run"] = [5000], 5000, 1000
    return [to for b in sq.plan(con, c2, SENDING, ignore_hours=True)[0].values() for _, to in b]


def stage(key, addr):
    v = con.execute("SELECT id FROM jobs WHERE status='new' AND send_status IS NULL AND job_type='full' LIMIT 1").fetchone()["id"]
    con.execute("UPDATE jobs SET status='matched', send_status='approved', emails=?, resume_path='x', dedupe_key=? WHERE id=?",
                (json.dumps([addr]), key, v))
    return v


@test("Sending", "every configured account sends, each with its own CC")
def _():
    want = sorted(e.lower() for e in cfg["sending"]["send_from"])
    assert sorted(s["email"].lower() for s in SENDING) == want, [s["email"] for s in SENDING]
    for sender, cc in cfg["sending"]["cc"].items():
        assert mailer.cc_for(sender, cfg) == cc.lower() and cc.lower() != sender.lower()
    return f"{len(SENDING)} accounts: " + ", ".join(s["email"].split("@")[0] for s in SENDING)


@test("Sending", "an address already emailed is never planned again")
def _():
    sent = con.execute("SELECT to_email FROM jobs WHERE send_status='sent' LIMIT 1").fetchone()
    if not sent:
        return "SKIP nothing sent yet"
    stage("zzz|zzz", sent["to_email"].upper())
    try:
        assert sent["to_email"].lower() not in planned(load_config())
    finally:
        con.rollback()
    return f"{sent['to_email']} blocked"


@test("Sending", "the same company is not written to again within 30 days")
def _():
    r = con.execute("SELECT dedupe_key FROM jobs WHERE send_status='sent' LIMIT 1").fetchone()
    if not r:
        return "SKIP nothing sent yet"
    stage(r["dedupe_key"].split("|")[0] + "|anotherrole", "another-address@gmail.com")
    try:
        assert "another-address@gmail.com" not in planned(load_config())
    finally:
        con.rollback()


@test("Sending", "a new company with a new address is planned")
def _():
    stage("brandnewcompanyzz|engineer", "hiring@brandnewcompanyzz.example")
    try:
        assert "hiring@brandnewcompanyzz.example" in planned(load_config())
    finally:
        con.rollback()


@test("Sending", "the database itself refuses a second email to an address")
def _():
    sent = con.execute("SELECT to_email FROM jobs WHERE send_status='sent' LIMIT 1").fetchone()
    if not sent:
        return "SKIP nothing sent yet"
    v = con.execute("SELECT id FROM jobs WHERE status='new' LIMIT 1").fetchone()["id"]
    try:
        con.execute("UPDATE jobs SET send_status='sending', to_email=? WHERE id=?", (sent["to_email"].upper(), v))
        raise AssertionError("duplicate was accepted")
    except sqlite3.IntegrityError:
        return "refused"
    finally:
        con.rollback()


@test("Sending", "daily cap follows the warm-up week")
def _():
    cap, week = sq.cap_today(con, SENDING[0]["email"], cfg)
    caps = cfg["sending"]["warmup_caps"]
    assert cap == min(caps[min(week - 1, len(caps) - 1)], cfg["sending"]["daily_ceiling"]), (cap, week)
    return f"week {week}: {cap} per day, {sq.sent_today(con, SENDING[0]['email'])} sent today"


@test("Sending", "nothing is planned once today's cap is used")
def _():
    c2 = load_config()
    c2["sending"]["warmup_caps"] = [sq.sent_today(con, SENDING[0]["email"])]
    assert not any(sq.plan(con, c2, SENDING, ignore_hours=True)[0].values())


@test("Sending", "outside office hours nothing is sent")
def _():
    sunday_night = dt.datetime(2026, 10, 11, 22, 0)
    assert not sq.in_hours(cfg, sunday_night) and sq.in_hours(cfg, dt.datetime(2026, 10, 12, 10, 0))


@test("Sending", "lead emails stay within their own daily limit")
def _():
    con.execute("UPDATE jobs SET send_status='approved' WHERE job_type='lead' AND status='matched' AND send_status IS NULL")
    try:
        c2 = load_config()
        c2["sending"]["warmup_caps"], c2["sending"]["daily_ceiling"], c2["sending"]["per_run"] = [5000], 5000, 1000
        counts = [sum(r["job_type"] == "lead" for r, _ in batch) for batch in sq.plan(con, c2, SENDING, ignore_hours=True)[0].values()]
        n = max(counts)
        assert 0 < n <= c2["leads"]["per_day_per_account"], counts
    finally:
        con.rollback()
    return f"{n} leads planned, limit {cfg['leads']['per_day_per_account']}"


@test("Sending", "an account with too many bounces stops for the day")
def _():
    ids = [r["id"] for r in con.execute("SELECT id FROM jobs WHERE status='new' AND send_status IS NULL LIMIT 10")]
    try:
        for n, i in enumerate(ids):
            con.execute("UPDATE jobs SET send_status='sent', sender='bouncetest@gmail.com', to_email=?, sent_at=?, outcome=? WHERE id=?",
                        (f"bt{n}@bounce-test.example", db.now(), "bounced" if n < 2 else None, i))
        assert sq.bounce_stop(con, "bouncetest@gmail.com", cfg)
    finally:
        con.rollback()


def _held_elsewhere(lock_class):
    """Start another process that holds the lock, as a running send or loop would."""
    import subprocess
    code = (f"import sys; sys.path.insert(0, {str(ROOT)!r}); from jobbot import sendqueue as sq; "
            f"l = sq.{lock_class}().__enter__(); print('held', flush=True); sys.stdin.read()")
    p = subprocess.Popen([sys.executable, "-c", code], stdin=subprocess.PIPE, stdout=subprocess.PIPE, text=True)
    assert p.stdout.readline().strip() == "held", "the helper process could not take the lock"
    return p


def _kill_tree(p):
    """Kill the helper and the interpreter a venv's python.exe starts under it, as a crash would."""
    if os.name == "nt":
        import subprocess
        subprocess.run(["taskkill", "/PID", str(p.pid), "/T", "/F"], capture_output=True)
    else:
        p.kill()
    p.wait()


@test("Sending", "two sends cannot run at the same time")
def _():
    other = _held_elsewhere("SendLock")
    try:
        held_by = sq.holder(sq.LOCK)   # a venv's python.exe starts the real interpreter as a child, so not other.pid
        assert held_by and held_by != os.getpid(), held_by
        try:
            with sq.SendLock():
                raise AssertionError("second send was allowed")
        except RuntimeError:
            pass
    finally:
        _kill_tree(other)
    assert sq.holder(sq.LOCK) is None, "a killed send left its lock behind"
    with sq.SendLock():
        assert sq.holder(sq.LOCK) == os.getpid()
    assert sq.holder(sq.LOCK) is None
    return "refused while another process sends; a killed send leaves no lock"


@test("Sending", "template-only drafts are held back when AI writing is required")
def _():
    from jobbot import cli_send
    if not (cfg.get("ai") or {}).get("require_ai_for_sending"):
        return "SKIP switched off in config: template emails are allowed"
    v = stage("aiTestCo|role", "jobs@gmail.com")
    con.execute("UPDATE jobs SET send_status=NULL, tailored_by='basic', score=99 WHERE id=?", (v,))
    try:
        con.commit()
        cli_send._auto_approve(con, load_config())
        state = con.execute("SELECT send_status FROM jobs WHERE id=?", (v,)).fetchone()[0]
        assert state is None, f"template-only draft was approved ({state})"
    finally:
        con.execute("UPDATE jobs SET status='new', send_status=NULL, emails='[]', resume_path=NULL, tailored_by=NULL, "
                    "dedupe_key='restored|' || id WHERE id=?", (v,))
        con.commit()


@test("Follow-up", "due only after 7 days, once, and never after a reply or bounce")
def _():
    from jobbot import followup as fu
    ids = [r["id"] for r in con.execute("SELECT id FROM jobs WHERE status='new' AND send_status IS NULL AND job_type='full' LIMIT 5")]
    old = (dt.date.today() - dt.timedelta(days=8)).isoformat() + "T10:00:00"
    recent = (dt.date.today() - dt.timedelta(days=2)).isoformat() + "T10:00:00"
    cases = [(old, None, None), (recent, None, None), (old, "replied", None), (old, "bounced", None), (old, None, db.now())]
    try:
        for i, (when, outcome, followed) in zip(ids, cases):
            con.execute("UPDATE jobs SET send_status='sent', sender=?, to_email=?, sent_at=?, outcome=?, followup_at=? WHERE id=?",
                        (SENDING[0]["email"], f"fu{i}@followup-test.example", when, outcome, followed, i))
        due = {r["id"] for r in fu.due(con, cfg)} & set(ids)
        assert due == {ids[0]}, f"expected only the first case, got {sorted(due)}"
        plan = fu.run(con, cfg, SENDING, {SENDING[0]["email"]: 5}, dry_run=True)
        assert any(job_id == ids[0] for job_id, _, _ in plan)
        assert not fu.run(con, cfg, SENDING, {SENDING[0]["email"]: 0}, dry_run=True), "sent although the daily cap is used"
    finally:
        con.rollback()
    return "1 of 5 cases due, as intended"


@test("Follow-up", "goes as a reply in the same thread, same recipient, no attachment")
def _():
    from jobbot import followup as fu
    r = con.execute("SELECT * FROM jobs WHERE send_status='sent' LIMIT 1").fetchone()
    if not r:
        return "SKIP nothing sent yet"
    sender = next(s for s in ACCOUNTS if s["email"] == r["sender"])
    m = fu.build(r, sender, master, cfg, "<original-id@mail.gmail.com>", r["email_subject"])
    assert m["To"] == r["to_email"] and m["In-Reply-To"] == "<original-id@mail.gmail.com>" and m["Subject"].startswith("Re: ")
    assert not m.is_multipart(), "follow-up carries an attachment"
    body = m.get_content()
    assert "follow up on my application" in body and "{" not in body and "gold card" not in body.lower()
    return m["Subject"][:70]


@test("Follow-up", "follow-ups count toward the account's daily cap")
def _():
    v = con.execute("SELECT id FROM jobs WHERE status='new' AND send_status IS NULL LIMIT 1").fetchone()["id"]
    before = sq.sent_today(con, "captest@gmail.com")
    try:
        con.execute("UPDATE jobs SET sender='captest@gmail.com', followup_at=? WHERE id=?", (db.now(), v))
        assert sq.sent_today(con, "captest@gmail.com") == before + 1
    finally:
        con.rollback()


@test("Pipeline", "only one repeating loop can run at a time")
def _():
    if sq.holder(sq.LOOP_LOCK):
        try:
            sq.LoopLock().__enter__()
            raise AssertionError("a second loop was allowed while one is running")
        except RuntimeError:
            return "a loop is running now; a second one is refused"
    other = _held_elsewhere("LoopLock")
    try:
        try:
            sq.LoopLock().__enter__()
            raise AssertionError("a second loop was allowed while one is running")
        except RuntimeError:
            pass
    finally:
        _kill_tree(other)
    with sq.LoopLock():
        assert sq.holder(sq.LOOP_LOCK) == os.getpid()
    assert sq.holder(sq.LOOP_LOCK) is None and not sq.pid_alive(999999)
    return "lock taken and released"


@test("Email", "message layout: one recipient, CC only your own address")
def _():
    m = mailer.build("upadhyaydivakar1991@gmail.com", "Divakar Upadhyay", "hr@gmail.com", "Subject", "Body", cc="divakarup1991@gmail.com")
    assert m["To"] == "hr@gmail.com" and m["Cc"] == "divakarup1991@gmail.com" and not m["Bcc"]
    assert not any(h.lower().startswith("x-") for h in m.keys()), list(m.keys())
    return "headers: " + ", ".join(m.keys())


@test("Email", "a foreign CC or two recipients are refused before sending")
def _():
    for to, cc in (("a@gmail.com, b@gmail.com", None), ("a@gmail.com", "stranger@gmail.com")):
        try:
            mailer.send({"email": "x", "password": "x", "smtp": "smtp.invalid"}, mailer.build("x@gmail.com", "N", to, "s", "b", cc=cc))
            raise AssertionError(f"accepted {to} / {cc}")
        except ValueError:
            pass


@test("Email", "both Gmail accounts can log in")
def _():
    for s in ACCOUNTS:
        mailer.check_login(s)
    return ", ".join(s["email"] for s in ACCOUNTS)


@test("Email", "a real email with PDF leaves the sending account (to your own inbox)")
def _():
    src = SENDING[0]
    dst = next(s["email"] for s in ACCOUNTS if s["email"] != src["email"])
    pdf = render.OUTPUT / "test" / "feature_test.pdf"
    mailer.send(src, mailer.build(src["email"], "Divakar Upadhyay", dst, "Job bot feature test " + db.now(),
                                  "This is the feature test of your job bot. Nothing was sent to any HR address.", pdf))
    return f"{src['email']} -> {dst}"


@test("Email", "inbox check for bounces and replies runs")
def _():
    b, r = sq.check_inbox(SENDING[0], con)
    return f"new bounces {b}, new replies {r}"


# ---------------------------------------------------------------- leads / collaboration
@test("Leads", "leads are generated directly from the tender feed")
def _():
    n = con.execute("SELECT COUNT(*) FROM jobs WHERE source='ted-tenders' AND status='matched' AND emails!='[]'").fetchone()[0]
    assert n > 0
    return f"{n} tender leads with a published contact email"


@test("Leads", "skip rules: no email, deadline too close, excluded country")
def _():
    soon = (dt.date.today() + dt.timedelta(days=2)).isoformat()
    later = (dt.date.today() + dt.timedelta(days=30)).isoformat()
    assert leads.skip_reason("", "Germany", later, cfg) == "no usable email published"
    assert leads.skip_reason("a@b.de", "Germany", soon, cfg) == "deadline passed or too close"
    assert "Taiwan" in leads.skip_reason("a@b.tw", "Taiwan", later, cfg)
    assert leads.skip_reason("a@b.de", "Germany", later, cfg) is None


@test("Leads", "lead email names the notice, has an opt-out line, no placeholders")
def _():
    subject, body = leads.write_email("City of Example", "New citizen portal", "https://example.org/n/1", "", master)
    assert "New citizen portal" in subject and "https://example.org/n/1" in body
    assert "will not write again" in body and "{" not in body + subject
    assert "gold card" not in body.lower()


@test("Leads", "lead emails always wait for your approval")
def _():
    from jobbot import cli_send
    lead = Job(source="test", ext_id="zz-test-lead", title="Tender: software", company="Test Agency", url="u",
               description="tender@lead-test.example", location="Remote", region="", remote=True, job_type="lead", posted_at="")
    try:
        db.insert_job(con, lead, {"track": "lead", "category": "x", "emails": ["tender@lead-test.example"], "score": 99,
                                  "notes": "", "skip": None, "dedupe_key": "leadtest|x"})
        con.execute("UPDATE jobs SET status='matched', job_type='lead', tailored_by='ai' WHERE ext_id='zz-test-lead'")
        auto = cfg | {"sending": cfg["sending"] | {"review": "auto"}}
        cli_send._auto_approve(con, auto)
        state = con.execute("SELECT send_status FROM jobs WHERE ext_id='zz-test-lead'").fetchone()[0]
        assert state is None, f"automatic review approved a lead email ({state})"
    finally:
        con.execute("DELETE FROM jobs WHERE ext_id='zz-test-lead'")
        con.execute("DELETE FROM postings WHERE ext_id='zz-test-lead'")
        con.commit()
    waiting = con.execute("SELECT COUNT(*) FROM jobs WHERE job_type='lead' AND status='matched' AND send_status IS NULL").fetchone()[0]
    return f"automatic review leaves leads alone; {waiting} waiting for you"


@test("Prospects", "city and district come from the address or its postal code")
def _():
    from jobbot import prospects as pr
    cases = {("台北市中山北路2段113號", "Zhongshan N. Rd.,Taipei City 104"): ("台北市", "中山"),
             ("臺北市內湖區瑞光路66號", ""): ("台北市", "內湖"), ("新竹縣竹北市台元一街8號", ""): ("新竹縣", "竹北"),
             ("新北巿土城區民生街九號", ""): ("新北市", "土城"), ("彰化市中山路3段359號", ""): ("彰化縣", "彰化"),
             ("新竹科學園區創新一路19號", ""): ("新竹市", "科學園區"),
             ("2F.,No.30,Zhongzheng Dist.,Taipei City 100028TAIPEI", "2F.,No.30,Zhongzheng Dist.,Taipei City 100028TAIPEI"): ("台北市", "中正")}
    bad = {k: pr.place(*k) for k, want in cases.items() if pr.place(*k) != want}
    assert not bad, bad
    assert pr.industry_code("IT") == "30" and pr.industry_code("資訊服務") == "30" and pr.city_matches("新竹縣", "Hsinchu")
    return f"{len(cases)} addresses"


@test("Prospects", "Taiwan email is Chinese then English, about one kind of work, mentions the Gold Card")
def _():
    from jobbot import prospects as pr
    row = {"industry": "30", "name_zh": "測試資訊股份有限公司", "search": "測試資訊", "short_name": "測試資訊"}
    for kind, o in pr.OFFER.items():
        subject, body = pr.write(row, master, "me@example.com", offer=kind)
        zh, _, en = body.partition("English version")
        assert "人力資源部 您好" in zh and "就業金卡" in zh and "敬祝 商祺" in zh, kind
        assert "Dear Hiring Team" in en and "Gold Card" in en, kind
        assert o["subj_zh"] in subject and o["subj_en"] in subject and "測試資訊" in subject, kind
        others = [x["subj_en"] for k, x in pr.OFFER.items() if k != kind]
        assert not any(x in subject for x in others), f"{kind} email offers more than one kind of work"
        assert not re.search(r"\{\w+\}", body.replace("me@example.com", "")) and "[" not in body, "placeholder or bracket left"
    return subject[:60]


@test("Prospects", "India email is English, has no Gold Card, names the city and offers one kind of work")
def _():
    from jobbot import prospects_india as pi
    row = {"short_name": "Test Softworks", "name_en": "Test Softworks", "search": "Test Softworks", "city": "Pune",
           "industry": "software industry"}
    subject, body = pi.write(row, master, "me@example.com")
    assert "Gold Card" not in body and "金卡" not in body and "relocating to Pune" in body
    assert "full-time" in subject and "freelance" not in subject and "Test Softworks" in subject
    subject2, body2 = pi.write(row | {"industry": "hospital"}, master, "me@example.com")
    assert "freelance" in subject2 and "Services I offer" in body2
    for b in (body, body2):
        assert not re.search(r"\{\w+\}", b) and "[" not in b, "placeholder or bracket left"
    assert pi.fit("software industry") == (1, "it") and pi.fit("pharmaceutical industry")[1] == "health"
    return subject[-40:]


@test("Prospects", "investor-relations addresses wait for you; general inboxes and company-domain addresses go")
def _():
    from jobbot import prospects as pr
    base = {"hr_email": None, "hr_email_where": None}
    assert pr.address({**base, "published_email": "ir@tsmc.com"})[1].startswith("CHECK")
    assert pr.address({**base, "published_email": "spokesman@acer.com"})[1].startswith("CHECK")
    assert not pr.address({**base, "published_email": "info@acer.com"})[1].startswith("CHECK")
    assert not pr.address({**base, "hr_email": "hr@x.com.tw", "hr_email_where": "site", "published_email": "ir@x.com.tw"})[1].startswith("CHECK")
    assert not pr._hr(["safe.workplace@fis.com.tw", "auditing@senao.com.tw"], {"safe.workplace@fis.com.tw"})


@test("Prospects", "speculative emails come after applications and stay within their own daily limit")
def _():
    c2 = load_config()
    c2["sending"]["warmup_caps"], c2["sending"]["daily_ceiling"], c2["sending"]["per_run"] = [5000], 5000, 1000
    batches = sq.plan(con, c2, SENDING, ignore_hours=True)[0].values()
    limit = (c2.get("prospects") or {}).get("per_day_per_account", 20)
    for batch in batches:
        kinds = [r["job_type"] for r, _ in batch]
        assert sum(k == "prospect" for k in kinds) <= limit, kinds
        if "prospect" in kinds:   # no job application after the first speculative email
            assert not {"full", "part"} & set(kinds[kinds.index("prospect"):]), kinds
    return f"limit {limit} per account"


@test("Collaboration", "contacts import creates drafts that wait for review")
def _():
    from openpyxl import Workbook
    path = render.OUTPUT / "test" / "collab_test.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook(); ws = wb.active
    ws.append(["Organisation", "Email", "Contact person", "Topic", "Website"])
    ws.append(["Feature Test Studio", "hello@gmail.com", "Sam", "joint web projects", "https://example.org"])
    wb.save(path)
    try:
        new, skipped, known = leads.import_contacts(str(path), cfg, master)
        r = con.execute("SELECT * FROM jobs WHERE source='collaboration' AND company='Feature Test Studio'").fetchone()
        assert new == 1 and r and r["track"] == "collaboration" and r["send_status"] is None
        assert r["email_subject"] == "Collaboration idea: joint web projects" and "Dear Sam," in r["email_body"]
    finally:
        con.execute("DELETE FROM jobs WHERE source='collaboration' AND company='Feature Test Studio'"); con.commit()
    return "draft created, then removed"


# ---------------------------------------------------------------- records
@test("Records", "tracker lists exactly the applications that were sent")
def _():
    import openpyxl
    path, n = rp.write_tracker(con)
    ws = openpyxl.load_workbook(path, read_only=True)["Applications sent"]
    in_db = con.execute("SELECT COUNT(*) FROM jobs WHERE send_status='sent'").fetchone()[0]
    assert ws.max_row - 1 == in_db == n, (ws.max_row - 1, in_db, n)
    return f"{n} rows in {path.name}"


@test("Records", "Excel job list has the columns you asked for")
def _():
    import openpyxl
    path, total, with_email = rp.write_jobs_book(con, cfg["max_age_days"])
    wb = openpyxl.load_workbook(path, read_only=True)
    head = [c.value for c in next(wb["All jobs"].iter_rows(max_row=1))]
    for col in ("Position", "Company", "Type", "Published", "Country", "City / town", "HR email", "Job link",
                "Company website", "Stage", "Job description"):
        assert col in head, f"missing column {col}"
    return f"{total} jobs, {with_email} with an HR email, sheets: {', '.join(wb.sheetnames)}"


@test("Records", "a sent job shows the status 'Email sent'")
def _():
    r = con.execute("SELECT * FROM jobs WHERE send_status='sent' AND outcome IS NULL LIMIT 1").fetchone()
    if not r:
        return "SKIP nothing sent yet"
    assert rp.stage(r) == "Email sent" and rp.delivery(r).startswith(("Sent", "Delivered"))


@test("Records", "date-wise CSV is written for today")
def _():
    jobs_path, n_jobs, emails_path, n_emails = rp.write_day(con, db.today())
    assert jobs_path.exists() and emails_path.exists()
    return f"{n_jobs} jobs, {n_emails} emails"


@test("Records", "cleanup deletes old PDFs and keeps the sent records")
def _():
    old = render.OUTPUT / "2026-01-01" / "999999"
    old.mkdir(parents=True, exist_ok=True)
    (old / "old.pdf").write_bytes(b"%PDF-1.4 test")
    before = con.execute("SELECT COUNT(*) FROM jobs WHERE send_status='sent'").fetchone()[0]
    pdfs, folders = cleanup.run(con)
    after = con.execute("SELECT COUNT(*) FROM jobs WHERE send_status='sent'").fetchone()[0]
    assert not old.exists() and pdfs >= 1 and before == after
    return f"{pdfs} PDFs deleted, {after} sent records kept"


# ---------------------------------------------------------------- API and admin page
PORT = 8799
import uvicorn  # noqa: E402
server = uvicorn.Server(uvicorn.Config(api.create_app(), host="127.0.0.1", port=PORT, log_level="warning", access_log=False))
threading.Thread(target=server.run, daemon=True).start()
while not server.started:
    time.sleep(0.05)


def _forget_test_accounts():
    """Remove everything earlier or current test runs created in the control database, and their folders."""
    import shutil
    for acc in control._all("SELECT * FROM accounts WHERE slug LIKE 'zz-test%'"):
        for table in ("schedules", "cycles", "audit", "payments", "tickets", "alerts", "suppressions", "applications",
                      "acct_settings", "acct_documents", "acct_prefs", "acct_meta", "acct_process_log"):   # the account's rows in the one database
            control._write(f"DELETE FROM {table} WHERE account_id=?", acc["id"])
        control._write("DELETE FROM tokens WHERE user_id IN (SELECT id FROM users WHERE account_id=?)", acc["id"])
        control._write("DELETE FROM users WHERE account_id=?", acc["id"])
        control._write("DELETE FROM accounts WHERE id=?", acc["id"])
        shutil.rmtree(control.home_of(acc), ignore_errors=True)
    mine = "SELECT id FROM users WHERE email LIKE 'zz-test%'"
    control._write(f"DELETE FROM audit WHERE user_id IN ({mine})")
    control._write(f"DELETE FROM cycles WHERE requested_by IN ({mine})")
    control._write(f"DELETE FROM schedules WHERE created_by IN ({mine})")
    control._write(f"DELETE FROM tokens WHERE user_id IN ({mine})")
    control._write("DELETE FROM users WHERE email LIKE 'zz-test%'")


# The tests sign in as a temporary owner of your own account, who is also the service operator.
_forget_test_accounts()
PASSWORD = "zz-test-" + os.urandom(9).hex()
control.user_create(1, "zz-test-owner@example.org", PASSWORD, "owner", operator=True)
TOKEN = control.login("zz-test-owner@example.org", PASSWORD)[0]


def call(path, method="GET", body=None, token="default"):
    token = TOKEN if token == "default" else token
    headers = {"Content-Type": "application/json"} if body is not None else {}
    if token:
        headers["Authorization"] = "Bearer " + token
    req = urllib.request.Request(f"http://127.0.0.1:{PORT}{path}", method=method,
                                 data=json.dumps(body).encode() if body is not None else None, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            raw = r.read()
            return r.status, (json.loads(raw) if "json" in r.headers.get("Content-Type", "") else raw.decode("utf-8"))
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read() or b"{}")


@test("API", "status, jobs, sent, sources, countries, processes answer")
def _():
    out = []
    for path in ("/status", "/jobs?has_email=1&limit=5", "/sent", "/sources", "/countries", "/processes", "/processes/log", "/runs"):
        code, body = call(path)
        assert code == 200 and body is not None, f"{path} -> {code}"
        out.append(path.split("?")[0])
    return ", ".join(out)


@test("API", "every listed route exists, is documented, and nothing else answers")
def _():
    spec = call("/openapi.json", token=None)[1]
    documented = {(m.upper(), path) for path, methods in spec["paths"].items() for m in methods}
    listed = {(m, path) for m, path, _, _ in api.ROUTES}
    assert listed == documented, f"only in the list: {listed - documented}; only in the docs: {documented - listed}"
    sample = {"job_id": "1", "name": "resume", "key_id": "0", "user_id": "0", "schedule_id": "0", "cycle_id": "0", "key": "x"}
    for method, path, _, _ in api.ROUTES:
        if method != "GET":
            continue
        url = path.format(**sample)
        code, body = call(url, token=None if (method, path) in api.OPEN else "default")
        unknown = isinstance(body, dict) and body.get("error") == "unknown path"
        assert code != 405 and not unknown, f"GET {path} is listed but nothing answers it ({code})"
    assert call("/no-such-route")[0] == 404 and call("/jobs", "DELETE")[0] == 405
    assert "swagger" in call("/docs", token=None)[1].lower()
    locked = [path for path, methods in spec["paths"].items() for m, op in methods.items() if op.get("security")]
    assert len(locked) > 40 and "/health" not in locked
    return f"{len(listed)} routes"


@test("API", "answers are not cached or framed, and an oversized request is refused")
def _():
    req = urllib.request.Request(f"http://127.0.0.1:{PORT}/status", headers={"Authorization": "Bearer " + TOKEN})
    with urllib.request.urlopen(req, timeout=30) as r:
        h = r.headers
    assert h["Cache-Control"] == "no-store" and h["X-Frame-Options"] == "DENY" and h["X-Content-Type-Options"] == "nosniff"
    try:   # refused with 413; on Windows the client may instead see the connection closed while it is still sending
        assert call("/candidates", "POST", {"url": "https://zz-too-big.example/", "region": "x" * (api.MAX_BODY + 10)})[0] == 413
    except (ConnectionError, urllib.error.URLError):
        pass
    assert not [c for c in sourcestore.candidates(q="zz-too-big")]
    assert call("/config", "PUT", ["not", "an", "object"])[0] == 400


@test("API", "admin panel, console and customer portal pages are served")
def _():
    code, page = call("/admin")
    assert code == 200 and "JobPilot Admin" in page, "admin panel"
    code, page = call("/console")
    assert code == 200 and "Job Bot Admin" in page and "Processes" in page and "v-settings" in page, "console"
    code, page = call("/app")
    assert code == 200 and "JobPilot" in page, "customer portal"


@test("API", "a process can be switched off and on through the API")
def _():
    assert call("/processes/leads/disable", "POST")[1]["enabled"] is False
    off = next(p for p in call("/processes")[1] if p["name"] == "leads")["enabled"]
    call("/processes/leads/enable", "POST")
    on = next(p for p in call("/processes")[1] if p["name"] == "leads")["enabled"]
    reset_setting("process.leads")
    assert off is False and on is True


@test("API", "a source can be switched off and on through the API")
def _():
    call("/sources/larajobs/disable", "POST")
    off = next(s for s in call("/sources")[1] if s["name"] == "larajobs")["enabled"]
    call("/sources/larajobs/enable", "POST")
    on = next(s for s in call("/sources")[1] if s["name"] == "larajobs")["enabled"]
    assert off is False and on is True


@test("API", "a source can be added, read, changed and deleted through the API")
def _():
    new = {"name": "zz-api-test", "adapter": "generic", "region": "Remote", "enabled": False,
           "home": "https://example.org/jobs", "params": {"kind": "feed", "url": "https://example.org/feed"}}
    try:
        code, made = call("/sources", "POST", new)
        assert code == 201 and made["name"] == "zz-api-test" and made["enabled"] is False, (code, made)
        assert call("/sources/zz-api-test")[1]["params"]["url"] == "https://example.org/feed"
        assert any(s["name"] == "zz-api-test" for s in call("/sources?enabled=0&q=zz-api")[1])
        assert not any(s["name"] == "zz-api-test" for s in call("/sources?enabled=1")[1])
        code, changed = call("/sources/zz-api-test", "PUT", {"region": "Japan", "enabled": True})
        assert code == 200 and changed["region"] == "Japan" and changed["enabled"] is True and changed["params"]["region"] == "Japan"
        assert call("/sources/zz-api-test", "PUT", {"name": "other"})[0] == 400
        assert call("/sources/zz-api-test", "PUT", {"params": {"kind": "feed", "url": "file:///etc/passwd"}})[0] == 400
    finally:
        gone = call("/sources/zz-api-test", "DELETE")
    assert gone[0] == 200 and call("/sources/zz-api-test")[0] == 404 and call("/sources/zz-api-test", "DELETE")[0] == 404


@test("API", "a bad source is refused and a built-in one cannot be deleted")
def _():
    before = len(call("/sources")[1])
    assert call("/sources", "POST", {"name": "zz-bad", "adapter": "nope", "params": {}})[0] == 400
    assert call("/sources", "POST", {"name": "Bad Name!", "adapter": "generic", "params": {"kind": "feed", "url": "https://example.org/f"}})[0] == 400
    assert call("/sources", "POST", {"name": "zz-bad", "adapter": "generic", "params": {"kind": "lever"}})[0] == 400
    assert call("/sources", "POST", {"url": "not a web address"})[0] == 400
    assert call("/sources/104", "DELETE")[0] == 409
    assert len(call("/sources")[1]) == before
    st = call("/status")[1]
    assert st["sources_total"] == before and st["sources_active"] == sum(s["enabled"] for s in call("/sources")[1])


@test("API", "candidate sites: list, filter, add")
def _():
    readable = call("/candidates?state=readable")[1]
    assert readable and all(c["state"] == "readable" for c in readable)
    url = "https://zz-candidate-test.example/jobs"
    try:
        assert call("/candidates", "POST", {"url": url, "region": "Remote"}) == (201, {"url": url, "added": True})
        assert call("/candidates", "POST", {"url": url, "region": "Remote"})[1]["added"] is False
        assert call("/candidates?state=untested&q=zz-candidate-test")[1][0]["state"] == "not tested yet"
        assert call("/candidates", "POST", {"url": "ftp://x"})[0] == 400
    finally:
        con.execute("DELETE FROM candidates WHERE url=?", (url,)); con.commit()
    return f"{len(readable)} readable sites"


@test("API", "settings list, and a value of the wrong kind is refused")
def _():
    rows = call("/settings")[1]
    assert rows and all({"key", "value", "default", "description", "changed"} <= set(r) for r in rows)
    assert call("/config", "PUT", {"loop.pause_minutes": "ten"})[0] == 400
    assert call("/config")[1]["effective"]["loop"]["pause_minutes"] == cfg["loop"]["pause_minutes"]
    return f"{len(rows)} settings"


@test("API", "profile documents: read, save, and refuse one that breaks the resume")
def _():
    names = [d["name"] for d in call("/profile")[1]]
    assert names == ["resume", "email-templates", "resume-template"], names
    resume, layout = call("/profile/resume")[1]["content"], call("/profile/resume-template")[1]["content"]
    assert resume == render.load_master() and "<" in layout
    before = render.to_html(render.build("taiwan"))
    assert call("/profile/resume", "PUT", {"content": resume})[0] == 200
    assert call("/profile/resume-template", "PUT", {"content": layout})[0] == 200
    assert call("/profile/email-templates", "PUT", {"content": call("/profile/email-templates")[1]["content"]})[0] == 200
    bad = [("resume", {k: v for k, v in resume.items() if k != "experience"}),      # a part is missing
           ("resume", resume | {"skills": "C#"}),                                  # wrong kind of value
           ("resume", resume | {"summary": resume["summary"] + " Taiwan Gold Card holder."}),   # would leak into resumes abroad
           ("resume-template", "{{ r.no_such_part.name }}"), ("resume-template", ""),
           ("email-templates", {"signature": "x"}), ("nope", {})]
    for name, content in bad:
        code, body = call(f"/profile/{name}", "PUT", {"content": content})
        assert code == 400 and body["error"], f"{name}: a broken version was saved"
    assert call("/profile/resume", "PUT", resume)[0] == 400 and call("/profile/nope")[0] == 404
    assert call("/profile/resume")[1]["content"] == resume and render.to_html(render.build("taiwan")) == before
    return f"{len(bad)} broken versions refused"


@test("API", "catalogue of known sites: list, add, remove")
def _():
    url = "https://zz-catalogue-test.example/jobs"
    everything = call("/catalogue")[1]
    assert everything and len(call("/catalogue?region=taiwan")[1]) < len(everything)
    try:
        assert call("/catalogue", "POST", {"url": url, "region": "Japan"})[0] == 201
        assert call("/catalogue", "POST", {"url": url, "region": "Japan"})[1]["added"] is False
        assert (url, "catalogue") in dc.candidates("Japan", use_search=False)
    finally:
        gone = call("/catalogue?url=" + urllib.parse.quote(url, safe=""), "DELETE")
    assert gone[0] == 200 and call("/catalogue?url=" + urllib.parse.quote(url, safe=""), "DELETE")[0] == 404
    assert len(call("/catalogue")[1]) == len(everything)
    return f"{len(everything)} sites"


@test("API", "settings: change, refuse a typo, refuse a protected one, reset")
def _():
    assert call("/config", "PUT", {"loop.pause_minutes": 12})[0] == 200
    assert call("/config")[1]["effective"]["loop"]["pause_minutes"] == 12
    assert call("/config", "PUT", {"loop.pause_minutez": 1})[0] == 400
    assert call("/config", "PUT", {"browser.channel": "x"})[0] == 400
    call("/config/loop.pause_minutes", "DELETE")
    assert call("/config")[1]["effective"]["loop"]["pause_minutes"] == cfg["loop"]["pause_minutes"]


@test("API", "passwords and keys are never exposed")
def _():
    from dotenv import dotenv_values
    secrets = [v for k, v in dotenv_values(DATA / ".env").items() if v and ("PASSWORD" in k or "KEY" in k)]
    blob = json.dumps([call(p)[1] for p in ("/config", "/status", "/sources", "/processes")])
    assert secrets and not any(s in blob or s.replace(" ", "") in blob for s in secrets)
    return f"{len(secrets)} secrets checked"


@test("API", "a sent application cannot be approved again")
def _():
    r = con.execute("SELECT id FROM jobs WHERE send_status='sent' LIMIT 1").fetchone()
    if not r:
        return "SKIP nothing sent yet"
    code, body = call(f"/jobs/{r['id']}/approve", "POST")
    assert code == 409 and body["changed"] is False


@test("API", "job detail carries description, email draft and interview prep")
def _():
    r = con.execute("SELECT id FROM jobs WHERE interview_prep IS NOT NULL LIMIT 1").fetchone()
    if not r:
        return "SKIP no interview prep written yet"
    body = call(f"/jobs/{r['id']}")[1]
    assert body["description"] and body["email_body"] and body["interview_prep"]


# ---------------------------------------------------------------- accounts, sign-in, schedules
@test("Accounts", "nothing is answered without signing in")
def _():
    paths = ("/", "/status", "/jobs", "/sent", "/sources", "/config", "/settings", "/profile/resume", "/cycles", "/me", "/audit")
    for path in paths:
        assert call(path, token=None)[0] == 401, f"{path} answered without a sign-in"
        assert call(path, token="jbs_made-up-token")[0] == 401, f"{path} answered to a made-up token"
    assert call("/sources/larajobs/disable", "POST", token=None)[0] == 401 and call("/cycles", "POST", token=None)[0] == 401
    assert call("/health", token=None)[0] == 200 and "JobPilot Admin" in call("/admin", token=None)[1]
    assert call("/setup", token=None)[1] == {"needed": False} and call("/setup", "POST", {}, token=None)[0] == 409
    return f"{len(paths)} paths closed, 3 open"


@test("Accounts", "sign-in: wrong password refused, repeated tries locked out, sign-out ends the token")
def _():
    assert call("/login", "POST", {"email": "zz-test-owner@example.org", "password": "wrong-password"}, token=None)[0] == 401
    code, body = call("/login", "POST", {"email": "zz-test-owner@example.org", "password": PASSWORD}, token=None)
    assert code == 200 and body["token"].startswith("jbs_") and "password_hash" not in body["user"]
    me = call("/me", token=body["token"])[1]
    assert me["user"]["email"] == "zz-test-owner@example.org" and me["account"]["id"] == 1 and me["not_ready"] is None
    assert call("/logout", "POST", token=body["token"])[0] == 200 and call("/me", token=body["token"])[0] == 401
    tries = [call("/login", "POST", {"email": "zz-test-nobody@example.org", "password": "x" * 12}, token=None)[0] for _ in range(7)]
    assert tries[:5] == [401] * 5 and tries[5:] == [429, 429], tries
    assert "scrypt$" in control._one("SELECT password_hash FROM users WHERE email='zz-test-owner@example.org'")["password_hash"]


@test("Accounts", "an API key works until it is revoked, and only its hash is stored")
def _():
    code, made = call("/keys", "POST", {"name": "zz test key"})
    assert code == 201 and made["key"].startswith("jbk_")
    assert call("/status", token=made["key"])[0] == 200
    listed = [k for k in call("/keys")[1] if k["name"] == "zz test key"]
    assert len(listed) == 1 and "key" not in listed[0] and made["key"] not in json.dumps(call("/keys")[1])
    assert not control._one("SELECT 1 FROM tokens WHERE token_hash=?", made["key"]), "the key itself is in the database"
    assert call(f"/keys/{listed[0]['id']}", "DELETE")[0] == 200 and call("/status", token=made["key"])[0] == 401


@test("Accounts", "roles: a viewer only reads, a member cannot change settings, an owner cannot lock themselves out")
def _():
    tokens = {}
    for role in ("viewer", "member"):
        code, u = call("/users", "POST", {"email": f"zz-test-{role}@example.org", "password": PASSWORD, "role": role})
        assert code == 201 and u["role"] == role and u["account_id"] == 1 and not u["operator"]
        tokens[role] = control.login(u["email"], PASSWORD)[0]
    assert call("/jobs?limit=1", token=tokens["viewer"])[0] == 200
    for method, path, body in (("POST", "/sources/larajobs/disable", None), ("POST", "/cycles", None), ("POST", "/keys", {"name": "x"}),
                               ("PUT", "/config", {"loop.pause_minutes": 11}), ("POST", "/schedules", {"kind": "interval", "every_minutes": 60})):
        assert call(path, method, body, token=tokens["viewer"])[0] == 403, f"a viewer could {method} {path}"
    assert call("/users", token=tokens["viewer"])[0] == 403 and call("/audit", token=tokens["member"])[0] == 403
    for method, path, body in (("PUT", "/config", {"loop.pause_minutes": 11}), ("PUT", "/profile/resume", {"content": {}}),
                               ("POST", "/schedules", {"kind": "interval", "every_minutes": 60}), ("POST", "/processes/leads/disable", None),
                               ("POST", "/users", {"email": "zz-test-x@example.org", "password": PASSWORD}), ("GET", "/accounts", None),
                               ("POST", "/accounts", {"name": "zz-test-nope", "email": "zz-test-y@example.org", "password": PASSWORD})):
        assert call(path, method, body, token=tokens["member"])[0] == 403, f"a member could {method} {path}"
    assert call("/sources/larajobs/disable", "POST", token=tokens["member"])[0] == 200
    assert call("/sources/larajobs/enable", "POST", token=tokens["member"])[0] == 200
    me = call("/me")[1]["user"]["id"]
    assert call(f"/users/{me}", "PUT", {"active": False})[0] == 400
    viewer = next(u for u in call("/users")[1] if u["email"] == "zz-test-viewer@example.org")
    assert call(f"/users/{viewer['id']}", "PUT", {"active": False})[1]["active"] == 0
    assert call("/jobs?limit=1", token=tokens["viewer"])[0] == 401, "a switched-off user is still signed in"
    assert call("/users", "POST", {"email": "zz-test-short@example.org", "password": "short"})[0] == 400


NEW = {}   # the second account made by the next test, used by the ones after it


@test("Accounts", "a new account has its own data and none of yours")
def _():
    code, made = call("/accounts", "POST", {"name": "zz-test customer", "timezone": "Europe/Berlin",
                                            "email": "zz-test-customer@example.org", "password": PASSWORD})
    assert code == 201 and made["account"]["home"] == "accounts/zz-test-customer" and made["owner"]["role"] == "owner", made
    NEW.update(account=made["account"], home=control.home_of(made["account"]),
               token=control.login("zz-test-customer@example.org", PASSWORD)[0])
    # the paywall is tested in the Portal area; here the account is on the internal (no charge) plan
    control._write("UPDATE accounts SET plan='internal' WHERE id=?", made["account"]["id"])
    assert NEW["home"].is_dir() and NEW["home"] != DATA and not (NEW["home"] / "jobbot.db").exists()   # files only; rows in the one database
    theirs = lambda path, *a: call(path, *a, token=NEW["token"])
    assert theirs("/jobs")[1]["total"] == 0 and theirs("/sent")[1] == [] and call("/jobs")[1]["total"] > 0
    assert theirs("/status")[1]["jobs_total"] == 0 and theirs("/status")[1]["accounts"] == []
    assert theirs("/me")[1]["not_ready"] == "the resume has not been filled in yet"
    assert theirs("/portal/resume")[1]["content"]["name"] == control.PLACEHOLDER_NAME
    # customers work only through the portal: the raw settings, documents and sources are closed to them
    for path, method in (("/config", "GET"), ("/profile/resume", "GET"), ("/sources", "GET"), ("/settings", "GET"),
                         ("/sources/larajobs/disable", "POST"), ("/config", "PUT"), ("/cycles", "POST"), ("/accounts", "GET")):
        code = theirs(path, method, {} if method == "PUT" else None)[0]
        assert code == 403, f"a customer reached {method} {path} ({code})"
    private = [master["name"], master["contact"]["email"], master["contact"]["phone"]] + cfg["sending"]["send_from"]
    blob = json.dumps([theirs(p)[1] for p in ("/status", "/portal/resume", "/portal/targets", "/portal/overview", "/portal/details",
                                              "/portal/cover-letter", "/portal/mailbox", "/me", "/keys", "/schedules", "/users")],
                      ensure_ascii=False)
    leaked = [x for x in private if x and x in blob]
    assert not leaked, f"your details are visible to the new account: {leaked}"
    assert "Gold Card" not in blob, "your Gold Card wording is in the new account's answers"
    # a change in one account does not touch the other
    targets = theirs("/portal/targets")[1]
    assert theirs("/portal/targets", "PUT", targets | {"daily_limit": 3})[0] == 200
    assert theirs("/portal/targets")[1].get("daily_limit") == 3
    assert call("/config")[1]["effective"]["loop"]["pause_minutes"] == cfg["loop"]["pause_minutes"]
    assert len(theirs("/users")[1]) == 1
    assert call("/accounts", "POST", {"name": "zz-test dup", "email": "zz-test-customer@example.org", "password": PASSWORD})[0] == 400
    assert not control._one("SELECT 1 FROM accounts WHERE slug='zz-test-dup'"), "an account was left behind by a refused request"
    with use_home(NEW["home"]):
        inside = [str(x) for x in (db.DB_FILE, render.OUTPUT, rp.REPORTS, sq.LOOP_LOCK)]
    assert all(x.startswith(str(NEW["home"])) for x in inside[1:]) and db.DB_FILE == DATA / "jobbot.db"   # one database for all
    return f"raw routes closed, 0 jobs, none of {len(private)} private values visible"

@test("Accounts", "schedules: saved in the account's time zone, bad ones refused, not visible to another account")
def _():
    theirs = lambda path, *a: call(path, *a, token=NEW["token"])
    code, daily = theirs("/schedules", "POST", {"name": "mornings", "kind": "daily", "at": "08:30", "days": [0, 1, 2, 3, 4]})
    assert code == 201 and daily["enabled"] is True, daily
    from zoneinfo import ZoneInfo
    local = dt.datetime.fromisoformat(daily["next_run_at"]).astimezone(ZoneInfo("Europe/Berlin"))
    assert (local.hour, local.minute) == (8, 30) and local.weekday() < 5 and local > dt.datetime.now(dt.timezone.utc)
    for bad in ({"kind": "interval", "every_minutes": 5}, {"kind": "daily", "at": "25:00", "days": [1]},
                {"kind": "daily", "at": "08:30", "days": []}, {"kind": "daily", "at": "08:30", "days": [9]}, {"kind": "hourly"}):
        assert theirs("/schedules", "POST", bad)[0] == 400, f"accepted {bad}"
    assert theirs(f"/schedules/{daily['id']}", "PUT", {"enabled": False})[1]["enabled"] is False
    assert call("/schedules")[1] == [] or all(x["id"] != daily["id"] for x in call("/schedules")[1])
    assert call(f"/schedules/{daily['id']}", "PUT", {"enabled": True})[0] == 404, "one account changed another's schedule"
    assert call(f"/schedules/{daily['id']}", "DELETE")[0] == 404 and theirs(f"/schedules/{daily['id']}", "DELETE")[0] == 200
    return "next run " + local.strftime("%a %H:%M %Z")


def _wait(account_id, cycle_id, want, seconds=40):
    end = time.time() + seconds
    while time.time() < end:
        scheduler.tick()
        run = scheduler.cycle(account_id, cycle_id)
        if run["status"] in want:
            return run
        time.sleep(0.5)
    return scheduler.cycle(account_id, cycle_id)


@test("Accounts", "a run is refused until the resume is filled in, then runs in the account's own folder")
def _():
    theirs = lambda path, *a: call(path, *a, token=NEW["token"])
    account_id, real = NEW["account"]["id"], scheduler.RUN_COMMAND
    scheduler.RUN_COMMAND = [sys.executable, "-c", "import os, pathlib; pathlib.Path(os.environ['JOBBOT_HOME'], 'ran.txt')"
                                                   ".write_text('x'); print('hello from the run')"]
    try:
        code, asked = theirs("/portal/bot/run", "POST")
        assert code == 201 and asked["queued"] and asked["cycle"]["status"] == "queued", (code, asked)
        skipped = _wait(account_id, asked["cycle"]["id"], ("skipped",))
        assert skipped["status"] == "skipped" and "resume" in skipped["note"], skipped
        resume = theirs("/portal/resume")[1]["content"] | {"name": "Test Customer"}
        assert theirs("/portal/resume", "PUT", {"content": resume})[0] == 200 and theirs("/me")[1]["not_ready"] is None
        # a schedule that has fallen due queues exactly one run
        every = theirs("/schedules", "POST", {"kind": "interval", "every_minutes": 120})[1]
        control._write("UPDATE schedules SET next_run_at='2000-01-01T00:00:00+00:00' WHERE id=?", every["id"])
        scheduler.tick()
        runs = [r for r in scheduler.cycles(account_id) if r["trigger"] == "schedule"]
        assert len(runs) == 1, f"{len(runs)} runs queued for one due schedule"
        assert theirs("/portal/bot/run", "POST")[1]["queued"] is False, "a second run was queued while one was open"
        done = _wait(account_id, runs[0]["id"], ("done", "failed"))
        assert done["status"] == "done" and done["exit_code"] == 0, done
        assert (NEW["home"] / "ran.txt").exists() and not (DATA / "ran.txt").exists(), "the run worked in the wrong folder"
        assert "hello from the run" in scheduler.log_tail(account_id, done["id"])
        assert call(f"/cycles/{done['id']}/log")[0] == 404, "the owner could read a customer's run log"
        after = next(x for x in theirs("/schedules")[1] if x["id"] == every["id"])
        assert after["next_run_at"] > control.now() and after["last_run_at"]
        theirs(f"/schedules/{every['id']}", "DELETE")
    finally:
        scheduler.RUN_COMMAND = real
    return "skipped, then done"

@test("Accounts", "a run that is going can be stopped from the API")
def _():
    theirs = lambda path, *a: call(path, *a, token=NEW["token"])
    real = scheduler.RUN_COMMAND
    scheduler.RUN_COMMAND = [sys.executable, "-c", "import time; time.sleep(120)"]
    try:
        code, asked = theirs("/portal/bot/run", "POST")
        assert code == 201, asked
        going = _wait(NEW["account"]["id"], asked["cycle"]["id"], ("running",))
        assert going["status"] == "running" and sq.pid_alive(going["pid"]), going
        assert theirs("/status")[1]["loop_running"] is True and call("/status")[1]["loop_running"] in (True, False)
        assert theirs("/portal/bot/off", "POST")[0] == 200
        time.sleep(1.5)
        ended = scheduler.cycle(NEW["account"]["id"], asked["cycle"]["id"])
        assert ended["status"] == "stopped" and not sq.pid_alive(going["pid"]), ended
    finally:
        scheduler.RUN_COMMAND = real

@test("Accounts", "every change through the API is recorded with who made it")
def _():
    log = call("/audit?limit=300")[1]
    mine = [a for a in log if a["email"] == "zz-test-owner@example.org"]
    assert any(a["action"] == "POST /users" and a["result"] == 201 for a in mine), "creating a user was not recorded"
    assert all(not a["action"].startswith("GET") for a in log) and PASSWORD not in json.dumps(log)
    theirs = call("/audit", token=NEW["token"])[1]
    assert theirs and all(a["email"] == "zz-test-customer@example.org" for a in theirs)
    return f"{len(mine)} changes by the test user"


@test("Pipeline", "a failed step is tried again, but never a step that sends email")
def _():
    from jobbot import cli_send
    calls, real_sleep, quiet = [], cli_send.time.sleep, cli_send.console.quiet
    cli_send.time.sleep, cli_send.console.quiet, cli_send.ITERATION["id"] = (lambda s: None), True, "zz-test-retry"

    def once_bad():
        calls.append(1)
        if len(calls) == 1:
            raise RuntimeError("site did not answer")
        return "fine"

    def always_bad():
        calls.append(1)
        raise RuntimeError("boom")
    try:
        assert cli_send._step("scrape", once_bad) == (True, "fine") and len(calls) == 2
        calls.clear()
        assert cli_send._step("scrape", always_bad)[0] is False and len(calls) == 1 + cfg["loop"]["step_retries"]
        for step in cli_send.NO_RETRY:
            calls.clear()
            assert cli_send._step(step, always_bad)[0] is False and len(calls) == 1, f"{step} was run twice"
    finally:
        cli_send.time.sleep, cli_send.console.quiet, cli_send.ITERATION["id"] = real_sleep, quiet, None
        con.execute("DELETE FROM process_log WHERE iteration='zz-test-retry'")
        con.commit()


# ---------------------------------------------------------------- browser and the whole run
@test("Browser", "Chrome (Playwright) opens a job page and finds the job on it")
def _():
    from jobbot import browser
    rows = con.execute("SELECT * FROM jobs WHERE source IN ('remoteok','remotive','hn-whoishiring') AND status='matched' "
                       "ORDER BY found_at DESC LIMIT 3").fetchall()
    # shown on screen unless TEST_HEADLESS=1 (you asked to watch the test Chrome)
    res = browser.check_jobs(rows, cfg["browser"]["channel"], headless=os.environ.get("TEST_HEADLESS") == "1")
    good = [r for r in res if r["title_found"] or r["company_found"]]
    assert good, [(r["http"], r["note"]) for r in res]
    return f"{len(good)} of {len(res)} pages confirmed"


@test("Pipeline", "the last full iteration ran every step without failure")
def _():
    # the newest iteration that ran to its end; one that is running right now, or that you stopped, is not judged
    it = con.execute("""SELECT iteration FROM process_log WHERE iteration IS NOT NULL AND step='report'
                        AND finished_at IS NOT NULL ORDER BY id DESC LIMIT 1""").fetchone()
    assert it, "no finished iteration recorded yet (run: jobbot cycle)"
    steps = con.execute("SELECT step, ok, note, finished_at FROM process_log WHERE iteration=? ORDER BY id", (it[0],)).fetchall()
    failed = [f"{s['step']}: {s['note']}" for s in steps if s["ok"] == 0]
    unfinished = [s["step"] for s in steps if s["finished_at"] is None]
    assert not failed, "; ".join(failed)
    assert not unfinished, f"still running or interrupted: {unfinished}"
    names = [s["step"] for s in steps]
    for need in ("clean", "scrape", "contacts", "match", "tailor", "inbox", "send", "report"):
        assert need in names, f"step {need} did not run"
    return f"iteration {it[0]}: {len(steps)} steps ok"


@test("Pipeline", "no duplicate addresses among everything sent")
def _():
    rows = [r[0].lower() for r in con.execute("SELECT to_email FROM jobs WHERE send_status='sent'")]
    assert len(rows) == len(set(rows))
    return f"{len(rows)} sent, all to different addresses"


# ---------------------------------------------------------------- operations
@test("Operations", "a backup is made, checked, and restores to the same data")
def _():
    from jobbot import ops
    info = ops.backup("test")
    try:
        ok, detail, counts = ops.verify(info["file"])
        assert ok and counts["postings"] > 0, detail
        drill = ops.restore_drill()
        assert drill["ok"] and drill["restored_counts"]["accounts"] == counts["accounts"], drill
        assert ops.backups()[0].name == Path(info["file"]).name
    finally:
        Path(info["file"]).unlink(missing_ok=True)
        Path(info["file"] + ".json").unlink(missing_ok=True)
    return f"{info['size_mb']} MB in {info['seconds']} s; restore drill ok"


@test("Operations", "an alert is recorded once, can be resolved, and is shown to the operator only")
def _():
    from jobbot import ops
    real, sent = ops._email_alert, []
    ops._email_alert = lambda to, kind, msg: sent.append(msg)
    key = "zz-test-alert-" + os.urandom(4).hex()
    try:
        assert ops.raise_alert("test", "zz-test: something broke", key=key) is True
        assert ops.raise_alert("test", "zz-test: something broke", key=key) is False, "the same alert was raised twice"
        alert = next(a for a in call("/operator/alerts")[1] if a["key"] == key)
        assert call("/operator/alerts", token=NEW["token"])[0] == 403
        assert call(f"/operator/alerts/{alert['id']}/resolve", "POST")[1]["resolved"] is True
        assert all(a["key"] != key for a in call("/operator/alerts")[1])
    finally:
        ops._email_alert = real
        control._write("DELETE FROM alerts WHERE key=?", key)
    return "raised once, resolved"


@test("Operations", "deep health check and metrics answer; metrics are closed to everyone but the operator")
def _():
    code, h = call("/health/deep", token=None)
    assert code in (200, 503) and h["database"] == "ok" and "disk_free_gb" in h, h
    assert set(call("/health", token=None)[1]) >= {"ok", "time"}
    code, text = call("/metrics")
    assert code == 200 and "jobbot_postings " in text and "jobbot_accounts " in text, text[:200]
    assert call("/metrics", token=None)[0] == 401 and call("/metrics", token=NEW["token"])[0] == 403
    return f"health {code}, {text.count('# TYPE')} metrics"


@test("Operations", "the service log is one JSON object per line")
def _():
    from jobbot import ops
    ops.event("zz-test event", account=0, note="test")
    line = (DATA / "logs" / "service.jsonl").read_text(encoding="utf-8").strip().splitlines()[-1]
    entry = json.loads(line)
    assert entry["msg"] == "zz-test event" and entry["account"] == 0 and entry["level"] == "INFO"
    return "ok"


@test("Security", "mailbox passwords are stored encrypted and still work")
def _():
    from jobbot import mailer, secrets_box
    sealed = secrets_box.seal("abcd efgh ijkl mnop")
    assert sealed.startswith("enc:") and "abcd" not in sealed and secrets_box.open_(sealed) == "abcd efgh ijkl mnop"
    assert secrets_box.open_("plain-old-value") == "plain-old-value" and secrets_box.seal(sealed) == sealed
    theirs = lambda path, *a: call(path, *a, token=NEW["token"])
    code, r = theirs("/portal/mailbox", "PUT", {"mailboxes": [{"email": "zz-test-box@example.org", "app_password": "abcdefghijklmnop"}]})
    assert code == 200, r
    raw = (NEW["home"] / ".env").read_text(encoding="utf-8")
    assert "abcdefghijklmnop" not in raw and "SENDER_1_APP_PASSWORD=enc:" in raw, "the password is stored in plain text"
    with use_home(NEW["home"]):
        assert mailer.senders()[0]["password"] == "abcdefghijklmnop"
    bad = theirs("/portal/mailbox", "PUT", {"mailboxes": [{"email": "zz-test-box@example.org", "app_password": "x",
                                                           "smtp": "smtp.example.org\nANTHROPIC_API_KEY=stolen"}]})
    assert bad[0] == 400 and "ANTHROPIC" not in (NEW["home"] / ".env").read_text(encoding="utf-8"), "a line could be injected"
    return "encrypted at rest, injection refused"


@test("Sending", "the do-not-contact list stops an email, for every account or for one")
def _():
    from jobbot import suppress
    t = Job(source="test", ext_id="zz-test-dnc", title="Senior .NET Developer", company="Dnc Test Co", url="u",
            description="hr@dnc-test.example", location="Remote", region="", remote=True, job_type="full", posted_at="")
    try:
        db.insert_job(con, t, {"track": "full", "category": "x", "emails": ["hr@dnc-test.example"], "score": 90,
                               "notes": "", "skip": None, "dedupe_key": "dnctest|x"})
        con.execute("UPDATE jobs SET status='matched', send_status='approved', resume_path='x' WHERE ext_id='zz-test-dnc'")
        con.commit()
        assert "hr@dnc-test.example" in [to for _, to in sq.sendable(con, cfg)]
        suppress.add("@dnc-test.example", "manual", account_id=999999)   # another account's own entry: not for this one
        assert "hr@dnc-test.example" in [to for _, to in sq.sendable(con, cfg)], "another account's entry blocked this one"
        suppress.add("hr@dnc-test.example", "unsubscribed")             # asked not to be written to: for everyone
        assert "hr@dnc-test.example" not in [to for _, to in sq.sendable(con, cfg)]
        assert suppress.OPT_OUT.search("Please remove me from your list") and not suppress.OPT_OUT.search("Thanks, we will review it")
        if not NEW.get("token"):
            return "service-wide and per-account entries obeyed (portal part runs with the Accounts area)"
        theirs = lambda path, *a: call(path, *a, token=NEW["token"])
        assert theirs("/portal/suppressions", "POST", {"address": "boss@zz-test.example"})[0] == 201
        mine = theirs("/portal/suppressions")[1]
        assert any(x["address"] == "boss@zz-test.example" for x in mine)
        assert theirs("/portal/suppressions/hr@dnc-test.example", "DELETE")[0] == 404, "a customer removed a service-wide entry"
        assert theirs("/portal/suppressions/boss@zz-test.example", "DELETE")[0] == 200
    finally:
        con.execute("DELETE FROM jobs WHERE ext_id='zz-test-dnc'")
        con.execute("DELETE FROM postings WHERE ext_id='zz-test-dnc'")
        con.commit()
        for a in ("@dnc-test.example", "hr@dnc-test.example", "boss@zz-test.example"):
            suppress.remove(a)
    return "service-wide and per-account entries obeyed"


@test("Security", "a mailbox is connected by signing in (OAuth): no password kept, token encrypted")
def _():
    if not NEW.get("token"):
        return "SKIP runs with the Accounts area"
    import base64 as b64
    from jobbot import mailer, oauth
    keys = {"JOBBOT_PUBLIC_URL": "https://jobpilot.example", "GOOGLE_OAUTH_CLIENT_ID": "zz-client",
            "GOOGLE_OAUTH_CLIENT_SECRET": "zz-secret"}
    saved = {k: os.environ.get(k) for k in keys}
    os.environ.update(keys)
    claims = b64.urlsafe_b64encode(json.dumps({"email": "zz-test-oauth@gmail.com"}).encode()).decode().rstrip("=")

    class Answer:
        status_code, headers = 200, {"content-type": "application/json"}

        def __init__(self, data):
            self._data = data

        def json(self):
            return self._data
    asked = []

    def fake_post(url, timeout=None, data=None):
        asked.append(data)
        if data["grant_type"] == "authorization_code":
            return Answer({"access_token": "at-1", "refresh_token": "rt-secret-1", "expires_in": 3600, "id_token": f"h.{claims}.s"})
        return Answer({"access_token": "at-2", "expires_in": 3600})
    real_post = oauth.requests.post
    oauth.requests.post = fake_post
    theirs = lambda path, *a: call(path, *a, token=NEW["token"])
    try:
        assert {"provider": "google", "name": "Gmail"} in theirs("/portal/mailbox/oauth")[1]["providers"]
        url = theirs("/portal/mailbox/oauth/start", "POST", {"provider": "google"})[1]["url"]
        q = urllib.parse.parse_qs(urllib.parse.urlparse(url).query)
        assert q["client_id"] == ["zz-client"] and q["code_challenge_method"] == ["S256"] and "mail.google.com" in q["scope"][0]
        assert q["redirect_uri"] == ["https://jobpilot.example/oauth/callback"]
        bad = call("/oauth/callback?state=made-up&code=x", token=None)
        assert bad[0] == 200 and "not connected" in bad[1].lower()
        page = call(f"/oauth/callback?state={q['state'][0]}&code=the-code", token=None)[1]
        assert "is connected" in page, page[:300]
        assert asked[0]["code_verifier"] and asked[0]["client_secret"] == "zz-secret"
        boxes = theirs("/portal/mailbox")[1]
        box = next(b for b in (boxes if isinstance(boxes, list) else boxes.get("mailboxes", [])) if b["email"] == "zz-test-oauth@gmail.com")
        assert box["connected_by"] == "google" and box["ready"] and not box["has_password"]
        raw = (NEW["home"] / ".env").read_text(encoding="utf-8")
        assert "rt-secret-1" not in raw and "_OAUTH_TOKEN=enc:" in raw, "the refresh token is stored in plain text"
        with use_home(NEW["home"]):
            s = next(x for x in mailer.senders() if x["email"] == "zz-test-oauth@gmail.com")
            oauth._access.clear()
            assert oauth.access_token(s) == "at-2" and asked[-1]["refresh_token"] == "rt-secret-1"
        assert oauth.xoauth2("a@b.c", "tok") == "user=a@b.c\x01auth=Bearer tok\x01\x01"
        assert call(f"/oauth/callback?state={q['state'][0]}&code=again", token=None)[1].count("not connected"), "a state was used twice"
        assert theirs("/portal/mailbox/zz-test-oauth@gmail.com", "DELETE")[0] == 200
        assert "rt-secret" not in (NEW["home"] / ".env").read_text(encoding="utf-8") and "OAUTH" not in (NEW["home"] / ".env").read_text(encoding="utf-8")
    finally:
        oauth.requests.post = real_post
        for k, v in saved.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
    return "signed in, token encrypted, renewed, disconnected"


@test("Portal", "Razorpay and Stripe: checkout opens their page, and only a correctly signed answer counts as paid")
def _():
    import hashlib
    import hmac as hm
    from jobbot import payments as pay
    keys = {"JOBBOT_PUBLIC_URL": "https://jobpilot.example", "RAZORPAY_KEY_ID": "rzp_test_x", "RAZORPAY_KEY_SECRET": "rz-secret",
            "RAZORPAY_WEBHOOK_SECRET": "rz-hook", "STRIPE_SECRET_KEY": "sk_test_x", "STRIPE_WEBHOOK_SECRET": "whsec_x"}
    saved = {k: os.environ.get(k) for k in keys}
    os.environ.update(keys)

    class Answer:
        status_code = 200

        def __init__(self, data):
            self._data = data

        def json(self):
            return self._data
    import requests as rq
    real_post, real_get, sent = rq.post, rq.get, []

    def fake_post(url, **kw):
        sent.append((url, kw))
        if "razorpay" in url:
            return Answer({"id": "plink_1", "short_url": "https://rzp.io/i/abc"})
        return Answer({"id": "cs_test_1", "url": "https://checkout.stripe.com/c/pay/cs_test_1"})
    rq.post = fake_post
    acc_id = control.account_create("zz-test payer")["id"]
    control._write("UPDATE accounts SET market='IN', plan='trial' WHERE id=?", acc_id)
    plan = control._one("SELECT * FROM plans WHERE market='IN' AND price_month>0 AND active=1 ORDER BY price_month")
    try:
        assert pay.ADAPTERS["razorpay"].ready() and pay.ADAPTERS["stripe"].ready()
        # Razorpay
        _, pid = control._write("INSERT INTO payments(account_id, plan, amount, currency, gateway, status, created_at) "
                                "VALUES(?,?,?,?,?,'created',?)", acc_id, plan["name"], plan["price_month"], "INR", "razorpay", control.now())
        nxt = pay.ADAPTERS["razorpay"].start(control._one("SELECT * FROM payments WHERE id=?", pid), {})
        assert nxt == {"mode": "redirect", "url": "https://rzp.io/i/abc"} and sent[-1][1]["json"]["amount"] == int(plan["price_month"] * 100)
        answer = {"razorpay_payment_link_id": "plink_1", "razorpay_payment_link_reference_id": f"jp-{pid}",
                  "razorpay_payment_link_status": "paid", "razorpay_payment_id": "pay_1"}
        forged = answer | {"razorpay_signature": "0" * 64}
        page = call(f"/payments/return?gateway=razorpay&payment={pid}&" + urllib.parse.urlencode(forged), token=None)[1]
        assert "not confirmed" in page.lower() and control._one("SELECT status FROM payments WHERE id=?", pid)["status"] == "created"
        sig = hm.new(b"rz-secret", "|".join(answer.values()).encode(), hashlib.sha256).hexdigest()
        page = call(f"/payments/return?gateway=razorpay&payment={pid}&" + urllib.parse.urlencode(answer | {"razorpay_signature": sig}), token=None)[1]
        assert "payment received" in page.lower(), page[:300]
        assert control._one("SELECT paid_until FROM accounts WHERE id=?", acc_id)["paid_until"] > control.now()
        # Stripe, confirmed by its signed webhook
        _, sid = control._write("INSERT INTO payments(account_id, plan, amount, currency, gateway, status, created_at) "
                                "VALUES(?,?,?,?,?,'created',?)", acc_id, plan["name"], plan["price_month"], "INR", "stripe", control.now())
        assert pay.ADAPTERS["stripe"].start(control._one("SELECT * FROM payments WHERE id=?", sid), {})["url"].startswith("https://checkout.stripe.com")
        event = json.dumps({"type": "checkout.session.completed", "data": {"object": {
            "id": "cs_test_1", "payment_status": "paid", "client_reference_id": str(sid), "payment_intent": "pi_1"}}}).encode()
        stamp = str(int(time.time()))
        good = hm.new(b"whsec_x", stamp.encode() + b"." + event, hashlib.sha256).hexdigest()

        def hook(sig_header):
            req = urllib.request.Request(f"http://127.0.0.1:{PORT}/webhooks/stripe", data=event, method="POST",
                                         headers={"Stripe-Signature": sig_header, "Content-Type": "application/json"})
            try:
                with urllib.request.urlopen(req, timeout=30) as r:
                    return r.status
            except urllib.error.HTTPError as e:
                return e.code
        assert hook(f"t={stamp},v1={'0' * 64}") == 400 and control._one("SELECT status FROM payments WHERE id=?", sid)["status"] == "created"
        assert hook(f"t={stamp},v1={good}") == 200
        assert control._one("SELECT status, reference FROM payments WHERE id=?", sid) == {"status": "paid", "reference": "pi_1"}
        assert hook(f"t={stamp},v1={good}") == 200, "a repeated webhook failed"
    finally:
        rq.post, rq.get = real_post, real_get
        control._write("DELETE FROM payments WHERE account_id=?", acc_id)
        control._write("UPDATE accounts SET slug='zz-test-payer-' || id WHERE id=?", acc_id)
        for k, v in saved.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
    return "forged answers refused; signed answers and webhooks recorded the payments"


@test("Portal", "a customer can download all their data and delete their account")
def _():
    code, reg = call("/register", "POST", {"name": "zz-test delete me", "email": "zz-test-delete@example.org", "password": PASSWORD,
                                           "market": "IN"}, token=None)
    assert code in (200, 201), reg
    me = lambda path, *a: call(path, *a, token=reg["token"])
    me("/portal/details", "PUT", {"full_name": "Zz Delete", "city": "Pune"})
    data = me("/portal/export")[1]
    assert data["account"]["id"] == reg["account"]["id"] and data["users"][0]["email"] == "zz-test-delete@example.org"
    assert "password_hash" not in json.dumps(data) and "Zz Delete" in json.dumps(data)
    assert me("/portal/account/delete", "POST", {"password": "wrong-password-1", "confirm": "DELETE"})[0] == 400
    gone = me("/portal/account/delete", "POST", {"password": PASSWORD, "confirm": "DELETE"})
    assert gone[0] == 200 and gone[1]["deleted"], gone
    assert me("/portal/export")[0] == 401, "the deleted account's sign-in still works"
    left = control._one("SELECT * FROM accounts WHERE id=?", reg["account"]["id"])
    assert left["status"] == "deleted" and left["contact_email"] is None and left["name"] == "Deleted account"
    assert not control._one("SELECT 1 FROM users WHERE email='zz-test-delete@example.org'")
    assert not control._one("SELECT 1 FROM acct_prefs WHERE account_id=?", reg["account"]["id"])
    assert call("/portal/account/delete", "POST", {"password": PASSWORD, "confirm": "DELETE"})[0] == 400, "the owner's account could be deleted"
    control._write("UPDATE accounts SET slug='zz-test-deleted-' || id WHERE id=?", reg["account"]["id"])   # cleared on the next run
    return f"export had {len(data)} sections; deleted, payments kept: {gone[1]['payments_kept']}"


@test("Portal", "a new customer sets up for free, is locked until paying, and is unlocked by a test payment")
def _():
    code, reg = call("/register", "POST", {"name": "zz-test portal", "email": "zz-test-portal@example.org", "password": PASSWORD,
                                           "market": "IN", "timezone": "Asia/Kolkata"}, token=None)
    assert code in (200, 201), reg
    me = lambda path, *a: call(path, *a, token=reg["token"])
    assert me("/portal/bot/run", "POST")[0] == 402, "the bot ran on a free plan"
    assert me("/portal/details", "PUT", {"full_name": "Zz Test", "city": "Pune"})[0] == 200, "set-up is not free"
    resume = me("/portal/resume")[1]["content"] | {"name": "Zz Test"}
    assert me("/portal/resume", "PUT", {"content": resume})[0] == 200
    targets = me("/portal/targets")[1]
    assert me("/portal/targets", "PUT", targets | {"titles": ["Data Analyst"]})[0] == 200
    assert me("/portal/onboarding/complete", "POST")[0] == 200
    locked = me("/portal/details", "PUT", {"full_name": "Changed"})
    assert locked[0] == 402 and locked[1].get("payment_required"), "the profile could be changed after set-up without paying"
    plan = next(p for p in me("/portal/billing")[1]["plans"] if p["price_month"] > 0)
    out = me("/portal/checkout", "POST", {"plan": plan["name"]})
    assert out[0] in (200, 201), out
    pid = out[1]["payment"]["id"]
    assert me(f"/portal/payments/{pid}/confirm", "POST", {"result": "fail"})[1].get("paid") is not True
    assert me("/portal/details", "PUT", {"full_name": "Changed"})[0] == 402
    out = me("/portal/checkout", "POST", {"plan": plan["name"]})
    paid = me(f"/portal/payments/{out[1]['payment']['id']}/confirm", "POST", {"result": "success"})
    assert paid[0] == 200, paid
    assert me("/portal/details", "PUT", {"full_name": "Changed"})[0] == 200, "still locked after paying"
    assert me("/portal/access")[1]["paid"] is True
    t = me("/portal/tickets", "POST", {"kind": "support", "subject": "zz-test", "message": "help please"})
    assert t[0] in (200, 201), t
    return f"locked, failed payment kept it locked, paid {plan['name']} unlocked it"


# ---------------------------------------------------------------- report
server.should_exit = True
_forget_test_accounts()
counts = {s: sum(r[2] == s for r in RESULTS) for s in ("PASS", "FAIL", "SKIP")}
lines = ["# Job Bot Feature Test Report", "", f"Run: {db.now().replace('T', ' ')}", "",
         f"**{counts['PASS']} passed, {counts['FAIL']} failed, {counts['SKIP']} skipped** out of {len(RESULTS)} tests.", "",
         "Nothing was sent to an HR address during these tests. One test email went from the sending account to your own other address.", "",
         "| # | Area | Feature tested | Result | Detail |", "|---|---|---|---|---|"]
for n, (area, name, state, detail, secs) in enumerate(RESULTS, 1):
    lines.append(f"| {n} | {area} | {name} | {state} | {detail.replace('|', '/')} |")
if not ONLY:
    (ROOT / "TEST-REPORT.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
print(f"\n{counts['PASS']} passed, {counts['FAIL']} failed, {counts['SKIP']} skipped."
      + ("" if ONLY else f" Report: {ROOT / 'TEST-REPORT.md'}"))
sys.exit(1 if counts["FAIL"] else 0)
