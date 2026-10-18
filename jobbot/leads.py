"""Lead generation: organisations that published a web / software requirement together with a
contact email. Leads come from two places:

  - directly from public tender feeds the bot reads itself (sources with track 'lead', e.g. EU TED)
  - from a lead sheet you import (`jobbot leads import <file>`)

Each lead becomes one short inquiry about that specific notice. No resume is attached.
Lead emails always wait for your approval and have their own small daily limit.
"""
import datetime as dt
import re

import openpyxl

from . import db, profilestore
from .filter import EMAIL_RE
from .sources.base import Job

HEADER_KEY = "Client (organisation)"


def _text(v):
    return str(v).strip() if v is not None else ""


def skip_reason(email, country, deadline, cfg):
    l_cfg = cfg.get("leads") or {}
    if not EMAIL_RE.fullmatch(email or ""):
        return "no usable email published"
    if country in (l_cfg.get("skip_countries") or {}):
        return l_cfg["skip_countries"][country]
    if country in (cfg.get("exclude_countries") or {}) or country in (cfg.get("remote_only_countries") or {}):
        return f"country excluded: {country}"
    try:
        if deadline and dt.date.fromisoformat(deadline[:10]) < dt.date.today() + dt.timedelta(days=l_cfg.get("min_days_left", 5)):
            return "deadline passed or too close"
    except ValueError:
        pass
    return None


def write_email(client, requirement, link, person, master, template="lead", buyer=""):
    t = profilestore.get("email-templates")
    c = master["contact"]
    title = re.sub(r"\s+", " ", requirement)
    short = title if len(title) <= 80 else title[:80].rsplit(" ", 1)[0] + "..."
    vals = {"title": short, "company": client, "greeting": f"Dear {person}," if person else "Dear Sir or Madam,",
            "link": link, "name": master["name"], "phone": c["phone"], "profile": c["links"][0],
            "base": "New Taipei City, Taiwan", "buyer": buyer or "a public body"}
    body = t[template]["body"].format(**vals).strip() + "\n\n" + t["lead"]["signature"].format(**vals).strip()
    return t[template]["subject"].format(**vals), re.sub(r"\n{3,}", "\n\n", body)


def add(con, cfg, master, *, source, client, requirement, email, country="", city="", link="", deadline="",
        posted="", kind="Web / software", note="", person="", website="", description="", track="lead", buyer=""):
    """Store one lead. Returns 'new', 'skipped' or 'known'."""
    email = (email or "").strip().lower()
    region = "taiwan" if country == "Taiwan" else "abroad"
    job = Job(source=source, ext_id=link or f"{client}|{requirement}"[:200], title=requirement, company=client, url=link,
              description=description or note, location=", ".join(x for x in (city, country) if x), region=region,
              remote=True, job_type="lead", track=track, posted_at=posted[:10])
    skip = skip_reason(email, country, deadline, cfg)
    norm = lambda s: re.sub(r"[\W_]+", "", s.lower())
    verdict = {"skip": skip, "score": 50, "notes": note, "emails": [email] if EMAIL_RE.fullmatch(email) else [],
               "dedupe_key": f"{norm(client)}|{norm(requirement)}", "track": track, "category": kind.replace("/", "-")}
    if not db.insert_job(con, job, verdict):
        return "known"
    if skip:
        return "skipped"
    subject, body = write_email(client, requirement, link, person, master, template=track, buyer=buyer)
    con.execute("""UPDATE jobs SET status='matched', matched_on=?, hr_name=?, deadline=?, email_subject=?, email_body=?,
                   tailored_by='template', email_source=?, contact_checked=1,
                   company_url=? WHERE source=? AND ext_id=?""",
                (db.today(), person, deadline[:10], subject, body,
                 "published in the official notice" if track in ("lead", "subcontract") else "from your contact sheet",
                 website or None, source, job.ext_id))
    return "new"


def add_from_job(con, cfg, master, job):
    """A lead produced by a source reader; details travel in job.extra."""
    x = job.extra
    return add(con, cfg, master, source=job.source, client=job.company, requirement=job.title, email=x.get("email", ""),
               country=x.get("country", ""), city=x.get("city", ""), link=job.url, deadline=x.get("deadline", ""),
               posted=job.posted_at, kind=x.get("kind", "Web / software"), note=x.get("note", ""),
               person=x.get("person", ""), website=x.get("website", ""), description=job.description,
               track=job.track or "lead", buyer=x.get("buyer", ""))


def read_sheet(path):
    """Rows of the first sheet that has the lead columns, as dicts keyed by the header row."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        for i, row in enumerate(rows[:15]):
            if row and HEADER_KEY in row and "Email" in row:
                head = [str(h).strip() if h else "" for h in row]
                return [dict(zip(head, r)) for r in rows[i + 1:] if r and any(r)]
    raise ValueError(f"No sheet with '{HEADER_KEY}' and 'Email' columns in {path}")


def import_file(path, cfg, master):
    """Returns (new, skipped, already_known)."""
    con = db.connect()
    counts = {"new": 0, "skipped": 0, "known": 0}
    for lead in read_sheet(path):
        client, req = _text(lead.get("Client (organisation)")), _text(lead.get("Requirement"))
        if not client or not req:
            continue
        counts[add(con, cfg, master, source="leads", client=client, requirement=req, email=_text(lead.get("Email")),
                   country=_text(lead.get("Country")), city=_text(lead.get("City / Region")),
                   link=_text(lead.get("Official notice / RFP link")), deadline=_text(lead.get("Deadline")),
                   posted=_text(lead.get("Posted date")), kind=_text(lead.get("Type")).split("(")[0].strip() or "Web / software",
                   note=_text(lead.get("Source")), person=_text(lead.get("Contact person")),
                   website=_text(lead.get("Client website")))] += 1
    con.commit()
    return counts["new"], counts["skipped"], counts["known"]


def import_contacts(path, cfg, master, track="collaboration"):
    """A simple sheet of people or organisations to write to. Needs a header row with 'Email' and
    one of 'Organisation' / 'Company' / 'Name'. Optional: 'Contact person', 'Topic', 'Website',
    'Country', 'Note'. Returns (new, skipped, already_known)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    con, counts = db.connect(), {"new": 0, "skipped": 0, "known": 0}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        for i, row in enumerate(rows[:15]):
            head = [str(h).strip() if h else "" for h in (row or [])]
            org = next((h for h in ("Organisation", "Organization", "Company", "Name") if h in head), None)
            if "Email" not in head or not org:
                continue
            for r in rows[i + 1:]:
                d = dict(zip(head, r))
                who, email = _text(d.get(org)), _text(d.get("Email"))
                if not who or not email:
                    continue
                counts[add(con, cfg, master, source=track, client=who, requirement=_text(d.get("Topic")) or "working together",
                           email=email, country=_text(d.get("Country")), link=_text(d.get("Website")),
                           person=_text(d.get("Contact person")), website=_text(d.get("Website")),
                           note=_text(d.get("Note")), kind=track.title(), track=track)] += 1
            break
    con.commit()
    return counts["new"], counts["skipped"], counts["known"]
